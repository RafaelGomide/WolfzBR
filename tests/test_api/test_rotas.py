"""Testes de integração das rotas (F1.G3–F1.G8)."""

import json
from pathlib import Path

import polars as pl
import pytest
from apoio_api import CORPUS, PORTA, TOKEN, XLSX, cliente, enviar

import wolfz
from wolfz.api.security import Settings
from wolfz.core.recipe import load_recipe, recipe_from_dict
from wolfz.core.workspace import Workspace
from wolfz.export.m import m_text

# ---------------------------------------------------------------------------
# Projetos
# ---------------------------------------------------------------------------


def test_criar_projeto_detecta_e_grava_a_receita_inicial(client, settings, projeto):
    assert projeto["name"] == "receita_municipal.xlsx"
    assert projeto["source"]["path"] is None and projeto["source"]["uploaded"] is True
    assert projeto["detection"]["cabecalho"] == [4, 4]
    assert projeto["detection"]["dados"] == [5, 16]
    passos = [p["op"] for p in projeto["recipe"]["steps"]]
    assert passos[:2] == ["select_range", "header_rows"]
    assert projeto["etag"].startswith('"')
    pasta = settings.workspace.project_dir(projeto["id"])
    for arquivo in ["project.json", "raw.parquet", "structure.json", "detection.json", "recipe.wolfz.json"]:
        assert (pasta / arquivo).is_file(), arquivo
    assert not any((settings.workspace.root / "uploads").iterdir())  # o upload temporário saiu


def test_receita_inicial_roda_no_arquivo_original(projeto):
    recipe = recipe_from_dict(projeto["recipe"])
    df = wolfz.run(recipe, CORPUS / "receita_municipal" / "original.xlsx")
    assert df.columns[:2] == ["Código", "Descrição da Receita"]
    assert df.height == 12


def test_criar_projeto_com_aba_escolhida(client):
    r = enviar(client, "licitacoes_ods", "original.ods", sheet="Legenda")
    assert r.status_code == 201, r.text
    assert r.json()["source"]["sheet"] == "Legenda"
    assert r.json()["recipe"]["source"] == {"type": "ods", "sheet": "Legenda"}


def _abas(client, nome, arquivo):
    caminho = CORPUS / nome / arquivo
    with caminho.open("rb") as f:
        return client.post("/api/sheets", files={"file": (nome + caminho.suffix, f, XLSX)})


def test_listar_abas_sem_criar_projeto(client, settings):
    r = _abas(client, "licitacoes_ods", "original.ods")
    assert r.status_code == 200, r.text
    assert [a["name"] for a in r.json()] == ["Licitações", "Legenda"]
    assert all(a["hidden"] is False for a in r.json())
    assert client.get("/api/projects").json() == []  # nada foi criado
    assert list((settings.workspace.root / "uploads").iterdir()) == []  # temporário apagado


def test_listar_abas_de_csv_e_vazio(client):
    r = client.post("/api/sheets", files={"file": ("a.csv", b"a;b\n1;2\n", "text/csv")})
    assert r.status_code == 200 and r.json() == []


def test_listar_abas_de_arquivo_ruim(client):
    r = client.post("/api/sheets", files={"file": ("extrato.xls", b"<html></html>", "x")})
    assert r.status_code == 422 and r.json()["code"] == "extension_mismatch"


def test_criar_projeto_com_arquivo_ruim_e_problem_details(client):
    r = client.post("/api/projects", files={"file": ("extrato.xls", b"<html><table></table></html>", "x")})
    assert r.status_code == 422
    corpo = r.json()
    assert r.headers["content-type"].startswith("application/problem+json")
    assert corpo["code"] == "extension_mismatch"
    assert corpo["instance"] == "/api/projects"
    assert "HTML" in corpo["detail"]


def test_criar_projeto_sem_arquivo(client):
    r = client.post("/api/projects")
    assert r.status_code == 422 and "file" in r.json()["detail"]


def test_listar_projetos(client, projeto):
    lista = client.get("/api/projects").json()
    assert [p["id"] for p in lista] == [projeto["id"]]
    assert lista[0]["open_here"] is True and lista[0]["open_elsewhere"] is False
    assert lista[0]["rows"] == 19


def test_abrir_projeto_devolve_receita_e_etag(client, projeto):
    r = client.get(f"/api/projects/{projeto['id']}")
    assert r.status_code == 200
    assert r.headers["ETag"] == projeto["etag"] == r.json()["etag"]
    assert r.json()["recipe"] == projeto["recipe"]


@pytest.mark.parametrize("pid", ["01ARZ3NDEKTSV4RRFFQ69G5FAV", "..%2F..%2Fetc", "nao-e-ulid"])
def test_projeto_inexistente_404(client, pid):
    r = client.get(f"/api/projects/{pid}")
    assert r.status_code == 404
    assert r.json()["status"] == 404


def test_apagar_projeto(client, settings, projeto):
    pasta = settings.workspace.project_dir(projeto["id"])
    assert client.delete(f"/api/projects/{projeto['id']}").status_code == 204
    assert not pasta.exists()
    assert client.get(f"/api/projects/{projeto['id']}").status_code == 404


def test_estrutura(client, projeto):
    corpo = client.get(f"/api/projects/{projeto['id']}/structure").json()
    assert corpo["structure"]["sheet"] == "Receita"
    assert corpo["structure"]["merged"]
    assert corpo["detection"]["cabecalho"] == [4, 4]


def test_perfil(client, projeto):
    corpo = client.get(f"/api/projects/{projeto['id']}/profile").json()
    assert corpo["total_rows"] == 12
    nomes = [c["name"] for c in corpo["columns"]]
    assert "Código" in nomes
    raw = client.get(f"/api/projects/{projeto['id']}/profile", params={"step": "raw", "sample": 5}).json()
    assert raw["total_rows"] == 19 and raw["sampled_rows"] == 5 and raw["is_sample"] is True
    assert client.get(f"/api/projects/{projeto['id']}/profile", params={"step": "s99"}).status_code == 404


# ---------------------------------------------------------------------------
# Receita: ETag, If-Match, gravação atômica (F1.G6)
# ---------------------------------------------------------------------------


def _url(projeto, rota):
    return f"/api/projects/{projeto['id']}/{rota}"


def test_get_receita_com_etag(client, projeto):
    r = client.get(_url(projeto, "recipe"))
    assert r.headers["ETag"] == projeto["etag"]
    assert r.json() == projeto["recipe"]


def test_put_exige_if_match(client, projeto):
    r = client.put(_url(projeto, "recipe"), json=projeto["recipe"])
    assert r.status_code == 428


def test_put_com_etag_velha_e_409(client, projeto):
    r = client.put(_url(projeto, "recipe"), json=projeto["recipe"], headers={"If-Match": '"velha"'})
    assert r.status_code == 409
    assert r.json()["etag"] == projeto["etag"]


def test_put_salva_e_troca_a_etag(client, settings, projeto):
    nova = {**projeto["recipe"], "name": "receita editada"}
    r = client.put(_url(projeto, "recipe"), json=nova, headers={"If-Match": projeto["etag"]})
    assert r.status_code == 200, r.text
    etag = r.json()["etag"]
    assert etag != projeto["etag"] and r.headers["ETag"] == etag
    pasta = settings.workspace.project_dir(projeto["id"])
    assert load_recipe(pasta / "recipe.wolfz.json").name == "receita editada"
    # A ETag antiga não salva mais: duas abas não se sobrescrevem.
    assert (
        client.put(_url(projeto, "recipe"), json=nova, headers={"If-Match": projeto["etag"]}).status_code
        == 409
    )
    assert client.put(_url(projeto, "recipe"), json=nova, headers={"If-Match": etag}).status_code == 200


def test_passo_desativado_nao_roda_no_preview(client, projeto):
    receita = projeto["recipe"]
    assert [p["op"] for p in receita["steps"]][:2] == ["select_range", "header_rows"]
    passos = [{**p, "enabled": p["op"] != "header_rows"} for p in receita["steps"]]
    r = client.put(
        _url(projeto, "recipe"), json={**receita, "steps": passos}, headers={"If-Match": projeto["etag"]}
    )
    assert r.status_code == 200, r.text
    assert r.json()["recipe"]["steps"][1]["enabled"] is False
    assert "enabled" not in r.json()["recipe"]["steps"][0]
    colunas = [c["name"] for c in client.get(_url(projeto, "preview")).json()["columns"]]
    assert colunas[0] == "A"  # sem o cabeçalho, a grade continua com as letras


def test_put_receita_invalida_aponta_o_passo(client, projeto):
    ruim = {**projeto["recipe"], "steps": [{"id": "s1", "op": "parse_brl", "cols": []}]}
    r = client.put(_url(projeto, "recipe"), json=ruim, headers={"If-Match": projeto["etag"]})
    assert r.status_code == 422
    assert r.json()["title"] == "Receita inválida"
    assert any("s1" in e for e in r.json()["errors"])


def test_falha_na_gravacao_nao_corrompe_a_receita(settings, monkeypatch):
    def disco_cheio(self, alvo):
        raise OSError(28, "No space left on device")

    with cliente(settings, raise_server_exceptions=False) as client:
        projeto = enviar(client).json()
        pasta = settings.workspace.project_dir(projeto["id"])
        antes = (pasta / "recipe.wolfz.json").read_bytes()
        with monkeypatch.context() as m:
            m.setattr(Path, "replace", disco_cheio)
            nova = {**projeto["recipe"], "name": "não vai"}
            r = client.put(_url(projeto, "recipe"), json=nova, headers={"If-Match": projeto["etag"]})
    assert r.status_code == 500
    assert r.json()["detail"] == "Erro inesperado. Veja o log do WolfzBR."
    assert (pasta / "recipe.wolfz.json").read_bytes() == antes


# ---------------------------------------------------------------------------
# Preview
# ---------------------------------------------------------------------------


def test_preview_da_receita_inteira(client, projeto):
    r = client.get(_url(projeto, "preview"), params={"limit": 5})
    assert r.status_code == 200
    corpo = r.json()
    assert corpo["total_rows"] == 12 and len(corpo["rows"]) == 5 and corpo["offset"] == 0
    assert corpo["columns"][0]["name"] == "Código"


def test_preview_do_raw_e_de_um_passo(client, projeto):
    raw = client.get(_url(projeto, "preview"), params={"step": "raw", "limit": 3}).json()
    assert raw["total_rows"] == 19 and [c["name"] for c in raw["columns"]] == ["A", "B", "C", "D"]
    s1 = client.get(_url(projeto, "preview"), params={"step": "s1"}).json()
    assert s1["total_rows"] == 13  # cabeçalho + 12 linhas


def test_preview_colunas_escolhidas(client, projeto):
    corpo = client.get(_url(projeto, "preview"), params={"columns": "Código", "limit": 2}).json()
    assert [c["name"] for c in corpo["columns"]] == ["Código"]
    assert client.get(_url(projeto, "preview"), params={"columns": "nada"}).status_code == 422


@pytest.mark.parametrize(
    ("params", "status"), [({"step": "s99"}, 404), ({"limit": 1001}, 422), ({"offset": -1}, 422)]
)
def test_preview_parametros_invalidos(client, projeto, params, status):
    r = client.get(_url(projeto, "preview"), params=params)
    assert r.status_code == status
    assert r.headers["content-type"].startswith("application/problem+json")


def test_erro_de_passo_no_preview_tem_step_id(client, projeto):
    receita = {**projeto["recipe"]}
    receita["steps"] = [*receita["steps"], {"id": "s99", "op": "rename", "mapping": {"nao_existe": "x"}}]
    etag = client.put(_url(projeto, "recipe"), json=receita, headers={"If-Match": projeto["etag"]}).json()[
        "etag"
    ]
    r = client.get(_url(projeto, "preview"))
    assert r.status_code == 422
    corpo = r.json()
    assert corpo["step_id"] == "s99" and corpo["title"] == "Coluna não encontrada"
    assert "Código" in corpo["available_columns"]
    assert corpo["instance"].endswith("/preview")
    assert etag


def test_um_preview_por_projeto_por_vez(client, projeto):
    projetos = client.app.state.projetos
    lock = projetos.lock_preview(projeto["id"])
    assert projetos.lock_preview(projeto["id"]) is lock  # o mesmo para todas as requisições do projeto
    assert lock.acquire(blocking=False)
    lock.release()


def test_handlers_pesados_rodam_em_threadpool(client):
    """Funções síncronas: o FastAPI as executa fora do event loop."""
    import inspect

    from wolfz.api.rotas import router

    api = router(client.app.state.settings, client.app.state.projetos)
    rotas = {r.path: r for r in api.routes if "GET" in r.methods or "POST" in r.methods}
    for caminho in [
        "/api/projects",
        "/api/projects/{project_id}/preview",
        "/api/projects/{project_id}/exports",
    ]:
        assert not inspect.iscoroutinefunction(rotas[caminho].endpoint), caminho


# ---------------------------------------------------------------------------
# Exportação
# ---------------------------------------------------------------------------


def test_exportar_formatos(client, projeto, tmp_path):
    destino = tmp_path / "saida"
    pedido = {"formats": ["csv", "parquet", "powerquery", "python"], "output_dir": str(destino)}
    r = client.post(_url(projeto, "exports"), json=pedido)
    assert r.status_code == 200, r.text
    nomes = sorted(Path(f).name for f in r.json()["files"])
    assert nomes == [
        "receita_municipal.CaminhoArquivo.pq",
        "receita_municipal.csv",
        "receita_municipal.parquet",
        "receita_municipal.pq",
        "receita_municipal.py",
    ]
    assert r.json()["rows"] == 12
    assert pl.read_parquet(destino / "receita_municipal.parquet").height == 12
    consulta = (destino / "receita_municipal.pq").read_text(encoding="utf-8")
    assert "Excel.Workbook(File.Contents(CaminhoArquivo), null, true)" in consulta
    assert "// s1 select_range" in consulta
    parametro = (destino / "receita_municipal.CaminhoArquivo.pq").read_text(encoding="utf-8")
    assert "receita_municipal.xlsx" in parametro
    assert "IsParameterQuery = true" in parametro
    assert not [p for p in destino.iterdir() if p.name.startswith(".wolfz-export")]  # sem temporários


def test_exportar_xlsx_e_csv_para_excel(client, projeto, tmp_path):
    pedido = {
        "formats": ["xlsx", "csv"],
        "output_dir": str(tmp_path),
        "csv_separator": ";",
        "decimal_comma": True,
    }
    r = client.post(_url(projeto, "exports"), json=pedido)
    assert r.status_code == 200, r.text
    assert r.json()["formula_cells"] == 0 and r.json()["neutralized"] is False
    from openpyxl import load_workbook

    ws = load_workbook(tmp_path / "receita_municipal.xlsx").active
    assert ws.max_row == 13 and ws["A1"].value == "Código"
    cabecalho = (tmp_path / "receita_municipal.csv").read_bytes()
    assert cabecalho.startswith(b"\xef\xbb\xbfC") and b";" in cabecalho.splitlines()[0]


def test_exportar_conta_e_neutraliza_formulas(client, projeto, tmp_path):
    receita = projeto["recipe"]
    passos = [
        *receita["steps"],
        {"id": "s9", "op": "replace", "cols": ["Código"], "old": "1.1.1.2.50.0.0", "new": "=1+1"},
    ]
    r = client.put(
        _url(projeto, "recipe"), json={**receita, "steps": passos}, headers={"If-Match": projeto["etag"]}
    )
    assert r.status_code == 200, r.text
    pedido = {"formats": ["csv"], "output_dir": str(tmp_path)}
    r = client.post(_url(projeto, "exports"), json=pedido)
    assert r.json()["formula_cells"] == 1 and r.json()["neutralized"] is False
    r = client.post(_url(projeto, "exports"), json={**pedido, "overwrite": True, "neutralize_formulas": True})
    assert r.json()["neutralized"] is True
    assert "\n'=1+1," in (tmp_path / "receita_municipal.csv").read_text(encoding="utf-8-sig")


def test_exportar_que_nao_cabe_e_422(client, projeto, tmp_path, monkeypatch):
    from wolfz.export import dados

    monkeypatch.setattr(dados, "MAX_LINHAS_XLSX", 5)
    saida = tmp_path / "saida"
    r = client.post(_url(projeto, "exports"), json={"formats": ["csv", "xlsx"], "output_dir": str(saida)})
    assert r.status_code == 422 and "Excel aceita" in r.json()["detail"]
    assert list(saida.iterdir()) == []  # nem o CSV foi gravado


def test_exportar_sobre_arquivo_existente(client, projeto, tmp_path):
    pedido = {"formats": ["csv"], "output_dir": str(tmp_path)}
    assert client.post(_url(projeto, "exports"), json=pedido).status_code == 200
    r = client.post(_url(projeto, "exports"), json=pedido)
    assert r.status_code == 409 and r.json()["existing"]
    assert client.post(_url(projeto, "exports"), json={**pedido, "overwrite": True}).status_code == 200


def test_exportar_na_pasta_do_projeto(client, settings, projeto):
    r = client.post(_url(projeto, "exports"), json={"formats": ["python"]})
    pasta = settings.workspace.project_dir(projeto["id"])
    assert Path(r.json()["files"][0]).parent == (pasta / "exports").resolve()
    assert r.json()["rows"] is None  # só código: não precisou executar a receita


@pytest.mark.parametrize(
    ("pedido", "trecho"),
    [
        ({"formats": ["pbip"]}, "ainda não disponível"),
        ({"formats": ["csv"], "csv_separator": "|"}, "csv_separator"),
        ({"formats": ["csv"], "output_dir": "relativo/pasta"}, "absoluto"),
        ({"formats": []}, "formats"),
        ({"formats": ["docx"]}, "formats"),
    ],
)
def test_exportar_pedido_invalido(client, projeto, pedido, trecho):
    r = client.post(_url(projeto, "exports"), json=pedido)
    assert r.status_code == 422 and trecho in r.json()["detail"]


def test_exportar_tableau(client, projeto, tmp_path):
    tmp_path = tmp_path / "saida"
    pytest.importorskip("tableauhyperapi", reason="extra wolfzbr[tableau] não instalado")
    r = client.post(_url(projeto, "exports"), json={"formats": ["hyper"], "output_dir": str(tmp_path)})
    assert r.status_code == 200, r.text
    assert sorted(Path(f).name for f in r.json()["files"]) == [
        "receita_municipal.hyper",
        "receita_municipal.tds",
    ]
    assert r.json()["rows"] == 12
    assert 'dbname="receita_municipal.hyper"' in (tmp_path / "receita_municipal.tds").read_text(
        encoding="utf-8"
    )
    assert sorted(p.name for p in tmp_path.iterdir()) == ["receita_municipal.hyper", "receita_municipal.tds"]


def test_tableau_sem_o_extra(client, projeto, tmp_path, monkeypatch):
    tmp_path = tmp_path / "saida"
    from wolfz.api import rotas

    monkeypatch.setattr(rotas, "hyper_disponivel", lambda: False)
    formatos = {f["format"]: f for f in client.get("/api/export-formats").json()}
    assert formatos["hyper"] == {"format": "hyper", "available": False, "install": rotas.INSTALAR}
    assert formatos["csv"]["available"] is True and "pbip" not in formatos
    r = client.post(_url(projeto, "exports"), json={"formats": ["csv", "hyper"], "output_dir": str(tmp_path)})
    assert r.status_code == 422 and "wolfzbr[tableau]" in r.json()["detail"]
    assert not tmp_path.exists()


# ---------------------------------------------------------------------------
# Catálogo
# ---------------------------------------------------------------------------


def test_catalogo_de_operacoes(client):
    ops = {o["op"]: o for o in client.get("/api/ops").json()}
    # Os testes de receita registram operações extras ("teste_…"); as do WolfzBR são 40.
    assert len([o for o in ops if not o.startswith("teste_")]) == 40
    assert ops["parse_brl"]["schema"]["properties"]["cols"]["type"] == "array"
    assert ops["sort"]["full_checkpoint"] is True and ops["rename"]["full_checkpoint"] is False
    assert ops["select_range"]["examples"] and ops["trim"]["doc"]


# ---------------------------------------------------------------------------
# Lock por projeto (F1.G7)
# ---------------------------------------------------------------------------


def test_projeto_aberto_em_outra_instancia_e_423(settings, projeto):
    outra = Settings(workspace=Workspace(settings.workspace.root), port=PORTA, token=TOKEN)
    with cliente(outra) as c:
        r = c.get(f"/api/projects/{projeto['id']}")
        assert r.status_code == 423
        assert f"http://127.0.0.1:{PORTA}" in r.json()["detail"]
        assert c.delete(f"/api/projects/{projeto['id']}").status_code == 423
        lista = c.get("/api/projects").json()
        assert lista[0]["open_elsewhere"] is True and lista[0]["open_here"] is False


def test_lock_e_liberado_ao_encerrar(settings):
    with cliente(settings) as c:
        pid = enviar(c).json()["id"]
        pasta = settings.workspace.project_dir(pid)
        assert json.loads((pasta / ".lock.json").read_text(encoding="utf-8"))["port"] == PORTA
    assert not (pasta / ".lock.json").exists()
    outra = Settings(workspace=Workspace(settings.workspace.root), port=PORTA, token=TOKEN)
    with cliente(outra) as c:
        assert c.get(f"/api/projects/{pid}").status_code == 200


def test_erro_inesperado_nao_vaza_detalhe(settings, monkeypatch):
    def explode(*args, **kwargs):
        raise RuntimeError("CPF 123.456.789-09 na mensagem")

    monkeypatch.setattr("wolfz.api.rotas.preview_project", explode)
    with cliente(settings, raise_server_exceptions=False) as client:
        projeto = enviar(client).json()
        r = client.get(_url(projeto, "preview"))
    assert r.status_code == 500
    assert "CPF" not in r.text
    assert r.json()["title"] == "Erro interno"


# ---------------------------------------------------------------------------
# Municípios: revisão e de-para (F2.C4, F2.C5)
# ---------------------------------------------------------------------------


@pytest.fixture
def projeto_municipios(client):
    r = enviar(client, "_fixtures", "municipios_digitados.csv")
    assert r.status_code == 201, r.text
    projeto = r.json()
    receita = projeto["recipe"]
    receita["steps"].append({"id": "s3", "op": "normalize_municipio", "col": "municipio", "uf_col": "uf"})
    r = client.put(_url(projeto, "recipe"), json=receita, headers={"If-Match": projeto["etag"]})
    assert r.status_code == 200, r.text
    return {**projeto, "recipe": r.json()["recipe"], "etag": r.json()["etag"]}


def test_revisar_municipios(client, projeto_municipios):
    r = client.get(_url(projeto_municipios, "steps/s3/municipios"))
    assert r.status_code == 200, r.text
    itens = r.json()
    assert itens[0]["valor"] == "Guarapary"
    assert itens[0]["status"] == "revisar"
    assert itens[0]["linhas"] == 2
    assert itens[0]["chave"] == "Guarapary|ES"
    assert itens[0]["sugestoes"][0]["nome"] == "Guarapari"
    assert {i["valor"]: i["status"] for i in itens} == {
        "Guarapary": "revisar",
        "Linhars": "fuzzy",
        "V. Velha": "abreviacao",
        "Serra": "exato",
        "Vitória": "exato",
    }


def test_revisar_municipios_depois_de_aceitar(client, projeto_municipios):
    receita = projeto_municipios["recipe"]
    receita["steps"][-1]["manual"] = {"Guarapary|ES": "3202405"}
    r = client.put(
        _url(projeto_municipios, "recipe"), json=receita, headers={"If-Match": projeto_municipios["etag"]}
    )
    assert r.status_code == 200, r.text
    itens = client.get(_url(projeto_municipios, "steps/s3/municipios")).json()
    guarapary = next(i for i in itens if i["valor"] == "Guarapary")
    assert (guarapary["status"], guarapary["nome_ibge"]) == ("manual", "Guarapari")


def test_revisar_passo_que_nao_e_de_municipio(client, projeto_municipios):
    assert client.get(_url(projeto_municipios, "steps/s1/municipios")).status_code == 422
    assert client.get(_url(projeto_municipios, "steps/s9/municipios")).status_code == 404


def test_buscar_municipios(client):
    r = client.get("/api/municipios", params={"q": "vila v", "uf": "ES"})
    assert r.status_code == 200
    achados = {(m["codigo_ibge"], m["nome"], m["uf"]) for m in r.json()}
    assert ("3205200", "Vila Velha", "ES") in achados
    assert all(uf == "ES" for _, _, uf in achados)
    assert client.get("/api/municipios", params={"q": ""}).status_code == 422


def test_exportar_com_depara(client, projeto_municipios, tmp_path):
    pedido = {"formats": ["powerquery", "python"], "output_dir": str(tmp_path)}
    r = client.post(_url(projeto_municipios, "exports"), json=pedido)
    assert r.status_code == 200, r.text
    nomes = sorted(Path(f).name for f in r.json()["files"])
    assert nomes == [
        "depara_municipios.csv",
        "fixtures.CaminhoArquivo.pq",
        "fixtures.CaminhoDeparaMunicipios.pq",
        "fixtures.pq",
        "fixtures.py",
    ]
    depara = pl.read_csv(tmp_path / "depara_municipios.csv", separator=";", infer_schema=False)
    assert depara.height == 5
    parametro = (tmp_path / "fixtures.CaminhoDeparaMunicipios.pq").read_text(encoding="utf-8")
    assert m_text(str(tmp_path / "depara_municipios.csv")) in parametro
    assert "CaminhoDeparaMunicipios" in (tmp_path / "fixtures.pq").read_text(encoding="utf-8")


# ---------------------------------------------------------------------------
# Várias tabelas na aba (F2.D3)
# ---------------------------------------------------------------------------


def test_aba_com_duas_tabelas(client):
    r = enviar(client, "anuario_tabelas_lado_a_lado")
    assert r.status_code == 201, r.text
    deteccao = r.json()["detection"]
    assert [t["colunas"] for t in deteccao["tabelas"]] == [["A", "C"], ["E", "G"]]
    assert deteccao["tabelas"][1]["passos"][0] == {"op": "select_range", "start": "E2", "end": "G10"}
    # A receita inicial usa a principal (empate: a primeira).
    assert r.json()["recipe"]["steps"][0]["start"] == "A2"
    estrutura = client.get(f"/api/projects/{r.json()['id']}/structure").json()
    assert len(estrutura["detection"]["tabelas"]) == 2


def test_aba_com_uma_tabela_nao_lista_tabelas(client, projeto):
    assert projeto["detection"].get("tabelas", []) == []


# ---------------------------------------------------------------------------
# Sugestões (F2.E)
# ---------------------------------------------------------------------------


def _sugestoes(client, projeto, regra=None):
    r = client.get(_url(projeto, "suggestions"))
    assert r.status_code == 200, r.text
    return [s for s in r.json() if regra is None or s["rule"] == regra]


def test_sugestoes(client, projeto):
    (sugestao,) = _sugestoes(client, projeto, "decimal_br")
    # Número do Excel chega como texto com ponto ('376357.13') e nenhuma
    # coluna tem vírgula: a tabela usa ponto decimal.
    assert sugestao["steps"] == [
        {"op": "parse_decimal_br", "cols": ["Previsão Atualizada (R$)", "Arrecadado (R$)"], "decimal": "."}
    ]
    assert sugestao["kind"] == "passo" and sugestao["confidence"] > 0.9
    # Verificações (F2.F2), por último: colunas 100% preenchidas viram
    # assert not_null; o código sem repetição vira assert unique.
    finais = _sugestoes(client, projeto)[-2:]
    assert {s["rule"] for s in finais} == {"verificar_preenchidas", "verificar_unicos"}
    assert all(s["steps"][0]["op"] == "assert" for s in finais)
    (unico,) = [s for s in finais if s["rule"] == "verificar_unicos"]
    assert unico["steps"] == [{"op": "assert", "check": "unique", "cols": ["Código"]}]


def test_rejeitar_sugestao_vale_para_o_projeto(client, settings, projeto):
    pedido = {"rule": "decimal_br", "cols": ["Arrecadado (R$)"]}
    r = client.post(_url(projeto, "suggestions/rejections"), json=pedido)
    assert r.status_code == 204, r.text
    client.post(_url(projeto, "suggestions/rejections"), json=pedido)  # repetir não duplica
    (sugestao,) = _sugestoes(client, projeto, "decimal_br")
    assert sugestao["cols"] == ["Previsão Atualizada (R$)"]
    gravado = json.loads(
        (settings.workspace.project_dir(projeto["id"]) / "suggestions.json").read_text(encoding="utf-8")
    )
    assert gravado == {"rejected": [{"rule": "decimal_br", "col": "Arrecadado (R$)"}]}

    client.post(_url(projeto, "suggestions/rejections"), json={"rule": "decimal_br"})
    assert _sugestoes(client, projeto, "decimal_br") == []


def test_aceitar_sugestao_tira_ela_da_lista(client, projeto):
    (sugestao,) = _sugestoes(client, projeto, "decimal_br")
    receita = projeto["recipe"]
    receita["steps"] += [{"id": f"s{len(receita['steps']) + 1}", **p} for p in sugestao["steps"]]
    r = client.put(_url(projeto, "recipe"), json=receita, headers={"If-Match": projeto["etag"]})
    assert r.status_code == 200, r.text
    assert _sugestoes(client, projeto, "decimal_br") == []


def test_verificacao_aceita_nao_volta(client, projeto):
    (verificacao,) = _sugestoes(client, projeto, "verificar_preenchidas")
    receita = projeto["recipe"]
    receita["steps"] += [{"id": f"s{len(receita['steps']) + 1}", **p} for p in verificacao["steps"]]
    r = client.put(_url(projeto, "recipe"), json=receita, headers={"If-Match": projeto["etag"]})
    assert r.status_code == 200, r.text
    assert _sugestoes(client, projeto, "verificar_preenchidas") == []


# ---------------------------------------------------------------------------
# Aba que muda de nome (F2.F4)
# ---------------------------------------------------------------------------


def test_importar_aba_com_mes_propoe_o_padrao(client, tmp_path):
    import openpyxl

    caminho = tmp_path / "relatorio.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "DEZEMBRO_2025"
    for linha in [["Município", "Valor"], ["Vitória", "10"], ["Serra", "20"]]:
        wb.active.append(linha)
    wb.save(caminho)
    with caminho.open("rb") as f:
        r = client.post("/api/projects", files={"file": ("relatorio.xlsx", f, XLSX)})
    assert r.status_code == 201, r.text
    fonte = r.json()["recipe"]["source"]
    assert fonte["sheet"] == "DEZEMBRO_2025"
    assert fonte["match"]["sheet_name_regex"].startswith("(?i)(?:")


def test_importar_aba_comum_nao_propoe_padrao(projeto):
    assert "match" not in projeto["recipe"]["source"]


def test_receita_de_versao_mais_nova_pede_para_atualizar(client, settings, projeto):
    """Projeto salvo por um WolfzBR mais novo (F2.G2): a tela do projeto mostra
    o título e o pedido de atualização, em vez de abrir a receita pela metade."""
    pasta = settings.workspace.project_dir(projeto["id"])
    caminho = pasta / "recipe.wolfz.json"
    dados = json.loads(caminho.read_text(encoding="utf-8"))
    caminho.write_text(json.dumps({**dados, "wolfz_version": "9.0"}), encoding="utf-8")
    r = client.get(f"/api/projects/{projeto['id']}")
    assert r.status_code == 422
    corpo = r.json()
    assert corpo["title"] == "Receita de uma versão mais nova do WolfzBR"
    assert "formato 9.0" in corpo["detail"]
    assert "pipx upgrade wolfzbr" in corpo["detail"]
    assert corpo["recipe_version"] == "9.0"
    assert corpo["supported_version"] == "0.1"
    # Salvar uma receita mais nova pela API também é recusado.
    caminho.write_text(json.dumps(dados), encoding="utf-8")
    r = client.put(
        f"/api/projects/{projeto['id']}/recipe",
        json={**projeto["recipe"], "wolfz_version": "9.0"},
        headers={"If-Match": projeto["etag"]},
    )
    assert r.status_code == 422
    assert r.json()["title"] == "Receita de uma versão mais nova do WolfzBR"
