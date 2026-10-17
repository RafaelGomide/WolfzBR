import datetime as dt

import openpyxl
import pytest

from wolfz.core.ingest import IngestError, list_sheets, read_sheet, read_sheet_raw
from wolfz.core.ingest.excel import _text


def test_xlsx_valores_como_texto_na_posicao_original(relatorio_xlsx):
    df = read_sheet_raw(relatorio_xlsx)
    # Grade começa em C3: colunas C..E e linha i = linha i + 1 do Excel.
    assert df.columns == ["C", "D", "E"]
    assert df.height == 10
    assert df.row(2) == ("Relatório de vendas", None, None)
    assert df.row(4) == ("CPF", "Valor", "Data")
    assert df.row(5) == ("01234567890", "1234.5", "2024-01-31 00:00:00")
    assert df.row(6) == ("00011122233", "10", "2024-02-01 13:45:00")
    assert df.row(9) == ("Fonte: sistema interno", None, None)
    assert all(str(t) == "String" for t in df.dtypes)


def test_xls_valores_como_texto_na_posicao_original(exemplo_xls):
    df = read_sheet_raw(exemplo_xls)
    assert df.columns == ["B", "C", "D"]
    assert df.row(1) == ("Relatório", None, None)
    assert df.row(3) == ("CPF", "Valor", "Data")
    # Mesmo formato de texto que o .xlsx produz.
    assert df.row(4) == ("01234567890", "1234.5", "2024-01-31 00:00:00")
    assert df.row(5) == ("00011122233", "10", "2024-02-01 00:00:00")


def test_ods_valores_como_texto(exemplo_ods):
    df = read_sheet_raw(exemplo_ods)
    assert df.columns == ["A", "B"]
    assert df.rows() == [
        (None, None),
        (None, "Relatório"),
        ("Município", "Valor"),
        ("São Paulo", "1234.5"),
        ("Vitória", "10"),
    ]


def test_aba_por_nome_e_por_indice(exemplo_ods):
    assert read_sheet_raw(exemplo_ods, "Outra").rows() == [("a",)]
    assert read_sheet_raw(exemplo_ods, 1).rows() == [("a",)]


def test_aba_inexistente(relatorio_xlsx):
    with pytest.raises(IngestError) as exc:
        read_sheet(relatorio_xlsx, "Nada")
    assert exc.value.code == "sheet_not_found"
    assert "'Vendas'" in exc.value.message


def test_primeira_aba_visivel_por_padrao(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Escondida"
    wb.active["A1"] = "não"
    wb.active.sheet_state = "hidden"
    wb.create_sheet("Visível")["A1"] = "sim"
    path = tmp_path / "a.xlsx"
    wb.save(path)
    assert read_sheet(path).sheet == "Visível"


def test_aba_vazia(tmp_path):
    wb = openpyxl.Workbook()
    path = tmp_path / "vazia.xlsx"
    wb.save(path)
    data = read_sheet(path)
    assert data.frame.shape == (0, 0)
    assert data.structure.sampled_rows == []


def test_listar_abas_xlsx(relatorio_xlsx):
    sheets = list_sheets(relatorio_xlsx)
    assert [(s.name, s.visibility, s.hidden) for s in sheets] == [
        ("Vendas", "visible", False),
        ("Oculta", "hidden", True),
        ("Interna", "very_hidden", True),
    ]
    assert (sheets[0].rows, sheets[0].columns) == (10, 5)
    assert sheets[0].to_dict()["hidden"] is False


def test_listar_abas_xls_e_ods(exemplo_xls, exemplo_ods):
    xls = list_sheets(exemplo_xls)
    assert [(s.name, s.hidden, s.rows, s.columns) for s in xls] == [
        ("Dados", False, 6, 4),
        ("Oculta", True, 1, 1),
    ]
    ods = list_sheets(exemplo_ods)
    assert [(s.name, s.rows) for s in ods] == [("Dados", 5), ("Outra", 1)]


def test_listar_abas_sem_dimensao(exemplo_xls):
    assert all(s.rows is None for s in list_sheets(exemplo_xls, dimensions=False))


def test_read_sheet_recusa_csv(tmp_path):
    path = tmp_path / "a.csv"
    path.write_text("a;b\n1;2\n", encoding="utf-8")
    with pytest.raises(ValueError):
        read_sheet(path)
    with pytest.raises(ValueError):
        read_sheet_raw(path)


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (None, None),
        ("", None),
        (True, "true"),
        (False, "false"),
        (12.0, "12"),
        (1234.5, "1234.5"),
        (1e22, "1e+22"),
        (7, "7"),
        (dt.datetime(2024, 1, 31, 13, 45, 7), "2024-01-31 13:45:07"),
        (dt.date(2024, 1, 31), "2024-01-31 00:00:00"),
        (dt.time(8, 30), "08:30:00"),
        (dt.timedelta(hours=1), "1:00:00"),
        ("  x ", "  x "),
    ],
)
def test_texto_do_calamine(value, expected):
    assert _text(value) == expected
