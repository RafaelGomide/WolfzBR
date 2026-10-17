"""Planilhas de teste geradas na hora (o ``.xls`` não tem gerador no
ambiente de desenvolvimento e fica em ``tests/corpus/_fixtures/``)."""

import datetime as dt
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl
import pytest
from openpyxl.styles import Font

from wolfz.core.workspace import Workspace

CORPUS = Path(__file__).parents[1] / "corpus" / "_fixtures"


@pytest.fixture
def exemplo_xls() -> Path:
    """Aba "Dados": título mesclado B2:D2, cabeçalho na linha 4, duas linhas
    de dado; aba "Oculta" escondida."""
    return CORPUS / "exemplo.xls"


@pytest.fixture
def relatorio_xlsx(tmp_path) -> Path:
    """Relatório típico: título mesclado em negrito, cabeçalho em negrito,
    dados, total com fórmula e nota de rodapé. Começa em C3."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Vendas"
    bold = Font(bold=True)
    ws["C3"] = "Relatório de vendas"
    ws["C3"].font = bold
    ws.merge_cells("C3:E3")
    for col, header in zip("CDE", ["CPF", "Valor", "Data"], strict=True):
        ws[f"{col}5"] = header
        ws[f"{col}5"].font = bold
    ws["C6"], ws["D6"], ws["E6"] = "01234567890", 1234.5, dt.datetime(2024, 1, 31)
    ws["C7"], ws["D7"], ws["E7"] = "00011122233", 10, dt.datetime(2024, 2, 1, 13, 45)
    ws["C8"], ws["D8"] = "Total", "=SUM(D6:D7)"
    ws["C8"].font = bold
    ws["C10"] = "Fonte: sistema interno"
    oculta = wb.create_sheet("Oculta")
    oculta["A1"] = "x"
    oculta.sheet_state = "hidden"
    muito = wb.create_sheet("Interna")
    muito["A1"] = "y"
    muito.sheet_state = "veryHidden"
    path = tmp_path / "relatorio.xlsx"
    wb.save(path)
    return path


_ODS_NS = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"'
)


def _ods_cell(value, span: int = 1) -> str:
    attrs = f' table:number-columns-spanned="{span}"' if span > 1 else ""
    if value is None:
        return "<table:table-cell/>"
    if isinstance(value, int | float):
        return (
            f'<table:table-cell office:value-type="float" office:value="{value}"{attrs}>'
            f"<text:p>{value}</text:p></table:table-cell>"
        )
    return (
        f'<table:table-cell office:value-type="string"{attrs}><text:p>{escape(str(value))}</text:p>'
        "</table:table-cell>"
    )


def write_ods(
    path: Path, sheets: dict[str, list[list]], merges: dict[tuple[str, int, int], int] | None = None
):
    """ODS mínimo. ``merges[(aba, linha, coluna)] = n`` mescla n colunas."""
    merges = merges or {}
    tables = []
    for name, rows in sheets.items():
        xml_rows = []
        for r, row in enumerate(rows):
            cells, c = [], 0
            while c < len(row):
                span = merges.get((name, r, c), 1)
                cells.append(_ods_cell(row[c], span))
                cells.extend("<table:covered-table-cell/>" for _ in range(span - 1))
                c += span
            xml_rows.append(f"<table:table-row>{''.join(cells)}</table:table-row>")
        tables.append(f'<table:table table:name="{escape(name)}">{"".join(xml_rows)}</table:table>')
    content = (
        f'<?xml version="1.0" encoding="UTF-8"?><office:document-content {_ODS_NS} office:version="1.2">'
        f"<office:body><office:spreadsheet>{''.join(tables)}</office:spreadsheet></office:body>"
        "</office:document-content>"
    )
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">'
        '<manifest:file-entry manifest:full-path="/" '
        'manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/>'
        '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
        "</manifest:manifest>"
    )
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet", zipfile.ZIP_STORED)
        zf.writestr("META-INF/manifest.xml", manifest, zipfile.ZIP_DEFLATED)
        zf.writestr("content.xml", content, zipfile.ZIP_DEFLATED)
    return path


@pytest.fixture
def exemplo_ods(tmp_path) -> Path:
    rows = [
        [None, None, None],
        [None, "Relatório", None],
        ["Município", "Valor", None],
        ["São Paulo", 1234.5, None],
        ["Vitória", 10, None],
    ]
    return write_ods(tmp_path / "exemplo.ods", {"Dados": rows, "Outra": [["a"]]}, {("Dados", 1, 1): 2})


@pytest.fixture
def workspace(tmp_path) -> Workspace:
    return Workspace(tmp_path / "home")
