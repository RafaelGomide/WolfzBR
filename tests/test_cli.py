import json
from pathlib import Path

import polars as pl
import pytest
from typer.testing import CliRunner

import wolfz
from wolfz.cli import app
from wolfz.core.recipe import recipe_from_dict, save_recipe

runner = CliRunner()


def test_version():
    result = runner.invoke(app, ["--version"])
    assert result.exit_code == 0
    assert result.output.strip() == f"wolfz {wolfz.__version__}"


def test_api_publica():
    assert wolfz.load_recipe is wolfz.core.recipe.load_recipe
    assert wolfz.run is wolfz.core.engine.run_file


def test_help_documenta_codigos_de_saida():
    result = runner.invoke(app, ["run", "--help"], terminal_width=200)
    assert result.exit_code == 0
    assert "0 ok" in result.output
    assert "1 um passo falhou" in result.output
    assert "2 receita inválida" in result.output
    assert "3 arquivo de entrada ilegível" in result.output


# ---------------------------------------------------------------------------
# wolfz run
# ---------------------------------------------------------------------------

PASSOS = [
    {"id": "s1", "op": "rename", "mapping": {"Vlr": "valor"}},
    {"id": "s2", "op": "drop_rows_where", "col": "regiao", "match": "starts_with", "value": "Total"},
    {"id": "s3", "op": "parse_brl", "cols": ["valor"]},
]


@pytest.fixture
def receita(tmp_path):
    path = tmp_path / "vendas.wolfz.json"
    data = {
        "wolfz_version": "0.1",
        "name": "vendas",
        "created_at": "2026-01-01T00:00:00Z",
        "source": {"type": "csv"},
        "steps": PASSOS,
    }
    save_recipe(recipe_from_dict(data), path)
    return path


@pytest.fixture
def entrada(tmp_path):
    path = tmp_path / "vendas_jan.csv"
    path.write_bytes("regiao;Vlr\nSul;R$ 1.234,56\nTotal;R$ 1.234,56\nNorte;(10,00)\n".encode("cp1252"))
    return path


def _run(*args):
    return runner.invoke(app, ["run", *map(str, args)])


def test_run_csv(receita, entrada, tmp_path):
    saida = tmp_path / "limpo.csv"
    result = _run(receita, entrada, "-o", saida)
    assert result.exit_code == 0, result.output
    assert f"2 linhas, 2 colunas -> {saida}" in result.output
    df = pl.read_csv(saida)
    assert df.to_dict(as_series=False) == {"regiao": ["Sul", "Norte"], "valor": [1234.56, -10.0]}


def test_run_parquet_em_pasta(receita, entrada, tmp_path):
    pasta = tmp_path / "saida"
    pasta.mkdir()
    result = _run(receita, entrada, "-o", pasta, "--format", "parquet")
    assert result.exit_code == 0, result.output
    df = pl.read_parquet(pasta / "vendas_jan.parquet")
    assert df.schema == {"regiao": pl.String, "valor": pl.Float64}


def test_run_pasta_nova_com_barra(receita, entrada, tmp_path):
    result = _run(receita, entrada, "-o", f"{tmp_path / 'nova'}/")
    assert result.exit_code == 0, result.output
    assert (tmp_path / "nova" / "vendas_jan.csv").exists()


def test_run_nao_sobrescreve_sem_flag(receita, entrada, tmp_path):
    saida = tmp_path / "limpo.csv"
    saida.write_text("antigo", encoding="utf-8")
    result = _run(receita, entrada, "-o", saida)
    assert result.exit_code == 2
    assert "A saída já existe" in result.output
    assert saida.read_text(encoding="utf-8") == "antigo"
    assert _run(receita, entrada, "-o", saida, "--overwrite").exit_code == 0
    assert saida.read_bytes().startswith(b"\xef\xbb\xbfregiao,valor")  # UTF-8 com BOM


def test_run_erro_de_passo_sai_com_1_e_nao_grava(receita, tmp_path):
    outro = tmp_path / "outro.csv"
    outro.write_text("regiao;Valor\nSul;1\n", encoding="utf-8")
    saida = tmp_path / "limpo.csv"
    result = _run(receita, outro, "-o", saida)
    assert result.exit_code == 1
    assert result.output.splitlines() == [
        "Coluna não encontrada: O passo s1 (rename) referencia a coluna 'Vlr', "
        "que não existe na tabela de entrada.",
        "Colunas disponíveis antes do passo s1: regiao, Valor",
        "Nada foi gravado.",
    ]
    assert not saida.exists()
    assert list(tmp_path.glob(".*.tmp")) == []


def test_mensagem_igual_a_da_api(receita, tmp_path):
    """O texto do terminal é o mesmo ``detail`` do Problem Details."""
    outro = tmp_path / "outro.csv"
    outro.write_text("regiao;Valor\nSul;1\n", encoding="utf-8")
    with pytest.raises(wolfz.StepError) as exc:
        wolfz.run(receita, outro)
    detail = exc.value.to_problem()["detail"]
    assert detail in _run(receita, outro, "-o", tmp_path / "x.csv").output


def test_run_receita_invalida_sai_com_2(entrada, tmp_path):
    receita = tmp_path / "ruim.wolfz.json"
    receita.write_text(
        json.dumps({"wolfz_version": "0.1", "steps": [{"id": "s1", "op": "xyz"}]}), encoding="utf-8"
    )
    result = _run(receita, entrada, "-o", tmp_path / "x.csv")
    assert result.exit_code == 2
    assert result.output.splitlines()[0] == f"Receita inválida: {receita}"
    assert "  s1: operação desconhecida: 'xyz'" in result.output.splitlines()


def test_run_receita_inexistente_sai_com_2(entrada, tmp_path):
    result = _run(tmp_path / "nao.wolfz.json", entrada, "-o", tmp_path / "x.csv")
    assert result.exit_code == 2
    assert "Receita não encontrada" in result.output


def test_run_json_quebrado_sai_com_2(entrada, tmp_path):
    receita = tmp_path / "quebrada.wolfz.json"
    receita.write_text('{"name": ', encoding="utf-8")
    result = _run(receita, entrada, "-o", tmp_path / "x.csv")
    assert result.exit_code == 2
    assert "JSON inválido na linha 1" in result.output


@pytest.mark.parametrize(
    ("nome", "conteudo", "mensagem"),
    [
        ("nao_existe.csv", None, "Arquivo não encontrado"),
        ("vendas.xlsx", b"PK\x03\x04" + b"lixo" * 20, "Arquivo corrompido"),
        ("vendas.xlsx", b"a;b\n1;2\n", "Extensão não bate com o conteúdo"),
        ("vendas.xls", b"<html><table></table></html>", "Extensão não bate com o conteúdo"),
    ],
)
def test_run_entrada_ilegivel_sai_com_3(receita, tmp_path, nome, conteudo, mensagem):
    arquivo = tmp_path / nome
    if conteudo is not None:
        arquivo.write_bytes(conteudo)
    result = _run(receita, arquivo, "-o", tmp_path / "x.csv")
    assert result.exit_code == 3
    assert mensagem in result.output


def test_run_sem_argumentos_e_uso_incorreto():
    assert runner.invoke(app, ["run"]).exit_code == 2


def test_schema():
    result = runner.invoke(app, ["schema"])
    assert result.exit_code == 0
    assert json.loads(result.output)["title"] == "Recipe"


def test_exemplo_de_ponta_a_ponta(tmp_path):
    """Critério de pronto da Fase 0: o exemplo do repositório roda pela CLI."""
    raiz = Path(__file__).resolve().parents[1]
    result = _run(raiz / "exemplo.wolfz.json", raiz / "tests" / "corpus" / "minimo.csv", "-o", f"{tmp_path}/")
    assert result.exit_code == 0, result.output
    assert pl.read_csv(tmp_path / "minimo.csv").to_dict(as_series=False) == {
        "regiao": ["Sul", "Sudeste", "Norte", "Centro-Oeste"],
        "valor": [1234.56, -10.0, 0.0, None],
        "obs": [None, "estorno", "sem venda no mês", "dado não informado"],
    }


# ---------------------------------------------------------------------------
# wolfz inspect
# ---------------------------------------------------------------------------


@pytest.fixture
def csv_brasileiro(tmp_path):
    path = tmp_path / "vendas.csv"
    linhas = ["Município;Valor;Data"] + [
        f"Cidade {i};R$ {i},00;{i % 28 + 1:02d}/01/2025" for i in range(1, 31)
    ]
    path.write_bytes("\n".join(linhas).encode("cp1252"))
    return path


def test_inspect_csv(csv_brasileiro):
    result = runner.invoke(app, ["inspect", str(csv_brasileiro)], env={"COLUMNS": "250"})
    assert result.exit_code == 0, result.output
    assert "vendas.csv (csv)" in result.output
    assert "30 linhas, 3 colunas" in result.output
    assert "amostra" not in result.output
    for trecho in ["Município", "moeda", "data", "01/01/2025", "28/01/2025"]:
        assert trecho in result.output


def test_inspect_amostra(csv_brasileiro):
    result = runner.invoke(app, ["inspect", str(csv_brasileiro), "--sample", "10"], env={"COLUMNS": "250"})
    assert result.exit_code == 0
    assert "Perfil numa amostra de 10 linhas" in result.output


def test_inspect_json(csv_brasileiro):
    result = runner.invoke(app, ["inspect", str(csv_brasileiro), "--json"])
    assert result.exit_code == 0
    dados = json.loads(result.output)
    assert dados["formato"] == "csv"
    assert dados["linha_cabecalho"] == 1
    assert dados["total_rows"] == 30
    valor = dados["columns"][1]
    assert (valor["name"], valor["inferred_type"], valor["minimum"], valor["maximum"]) == (
        "Valor",
        "moeda",
        1.0,
        30.0,
    )


def test_inspect_planilha_com_aba_e_cabecalho(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active["A1"] = "Título"
    outra = wb.create_sheet("Oculta")
    outra.sheet_state = "hidden"
    outra["A1"] = "x"
    dados = wb.create_sheet("Dados")
    dados.append(["Relatório"])
    dados.append(["CPF", "Valor"])
    dados.append(["123.456.789-09", 10])
    path = tmp_path / "rel.xlsx"
    wb.save(path)
    result = runner.invoke(
        app, ["inspect", str(path), "--sheet", "Dados", "--header-row", "2"], terminal_width=250
    )
    assert result.exit_code == 0, result.output
    assert "Aba: Dados" in result.output
    assert "Oculta (oculta)" in result.output
    assert "Cabeçalho na linha 2" in result.output
    assert "cpf" in result.output


def test_inspect_arquivo_ilegivel(tmp_path):
    path = tmp_path / "falso.xls"
    path.write_bytes(b"<html><table><tr><td>1</td></tr></table></html>")
    result = runner.invoke(app, ["inspect", str(path)])
    assert result.exit_code == 3
    assert "Extensão não bate com o conteúdo" in result.output


def test_inspect_cabecalho_invalido(tmp_path):
    path = tmp_path / "a.csv"
    path.write_text("a;b\n1;2\n", encoding="utf-8")
    result = runner.invoke(app, ["inspect", str(path), "--header-row", "3"])
    assert result.exit_code == 2
    assert "Uso incorreto" in result.output


# ---------------------------------------------------------------------------
# Exportação de dados (F1.I)
# ---------------------------------------------------------------------------


def test_run_csv_para_excel_brasileiro(receita, entrada, tmp_path):
    saida = tmp_path / "limpo.csv"
    assert _run(receita, entrada, "-o", saida, "--sep", ";", "--decimal-comma").exit_code == 0
    linhas = saida.read_text(encoding="utf-8-sig").splitlines()
    assert linhas == ["regiao;valor", "Sul;1234,56", "Norte;-10,0"]


def test_run_xlsx(receita, entrada, tmp_path):
    from openpyxl import load_workbook

    result = _run(receita, entrada, "-o", f"{tmp_path}/", "--format", "xlsx")
    assert result.exit_code == 0, result.output
    ws = load_workbook(tmp_path / "vendas_jan.xlsx")["vendas"]
    assert [[c.value for c in linha] for linha in ws.iter_rows()] == [
        ["regiao", "valor"],
        ["Sul", 1234.56],
        ["Norte", -10.0],
    ]


def test_run_avisa_e_neutraliza_formulas(receita, tmp_path):
    entrada = tmp_path / "perigo.csv"
    entrada.write_text("regiao;Vlr\n=HYPERLINK(1);R$ 1,00\n", encoding="utf-8")
    saida = tmp_path / "limpo.csv"
    result = _run(receita, entrada, "-o", saida)
    assert result.exit_code == 0 and "1 células de texto começam com =" in result.output
    assert "\n=HYPERLINK(1)," in saida.read_text(encoding="utf-8-sig")
    result = _run(receita, entrada, "-o", saida, "--overwrite", "--neutralize-formulas")
    assert result.exit_code == 0 and "Atenção" not in result.output
    assert "\n'=HYPERLINK(1)," in saida.read_text(encoding="utf-8-sig")


# ---------------------------------------------------------------------------
# Verificação na reaplicação (F2.F1) e lote (F2.F3)
# ---------------------------------------------------------------------------


@pytest.fixture
def receita_com_verificacao(tmp_path):
    path = tmp_path / "vendas_verificada.wolfz.json"
    data = {
        "wolfz_version": "0.1",
        "name": "vendas",
        "created_at": "2026-01-01T00:00:00Z",
        "source": {"type": "csv"},
        "steps": [
            *PASSOS,
            {"id": "s4", "op": "assert", "check": "not_null", "cols": ["valor"]},
            {"id": "s5", "op": "assert", "check": "row_count_min", "min_rows": 2},
        ],
    }
    save_recipe(recipe_from_dict(data), path)
    return path


def _csv(path, linhas):
    path.write_bytes(("regiao;Vlr\n" + "".join(f"{r};{v}\n" for r, v in linhas)).encode("cp1252"))
    return path


def test_verificacao_reprovada_sai_com_1_e_nao_grava(receita_com_verificacao, tmp_path):
    entrada = _csv(tmp_path / "mes2.csv", [("Sul", "R$ 1,00"), ("Norte", "")])
    saida = tmp_path / "saida.csv"
    result = _run(receita_com_verificacao, entrada, "-o", saida)
    assert result.exit_code == 1
    assert "Verificação reprovada" in result.stderr
    assert "'valor' tem 1 linhas vazia" in result.stderr
    assert not saida.exists()


def test_lote_com_relatorio(receita_com_verificacao, tmp_path):
    entradas = tmp_path / "entrada"
    entradas.mkdir()
    _csv(entradas / "jan.csv", [("Sul", "R$ 1,00"), ("Norte", "R$ 2,00")])
    _csv(entradas / "fev.csv", [("Sul", "R$ 3,00"), ("Norte", "")])  # verificação reprova
    _csv(entradas / "mar.csv", [("Sul", "R$ 4,00"), ("Leste", "R$ 5,00"), ("Oeste", "R$ 6,00")])
    saida, relatorio = tmp_path / "saida", tmp_path / "relatorio.json"
    # O padrão vai entre aspas: no PowerShell ninguém expande, a CLI expande.
    result = runner.invoke(
        app,
        [
            "run",
            str(receita_com_verificacao),
            str(entradas / "*.csv"),
            "-o",
            str(saida),
            "--report",
            str(relatorio),
        ],
        terminal_width=200,
    )
    assert result.exit_code == 1, result.output
    assert sorted(p.name for p in saida.iterdir()) == ["jan.csv", "mar.csv"]
    assert "2 de 3 arquivos" in result.stdout
    dados = json.loads(relatorio.read_text(encoding="utf-8"))
    por_arquivo = {Path(r["arquivo"]).name: r for r in dados["arquivos"]}
    assert por_arquivo["jan.csv"]["ok"] and por_arquivo["jan.csv"]["linhas"] == 2
    assert por_arquivo["mar.csv"]["linhas"] == 3
    assert not por_arquivo["fev.csv"]["ok"] and por_arquivo["fev.csv"]["codigo"] == 1
    assert "Verificação reprovada" in por_arquivo["fev.csv"]["erro"]


def test_lote_tudo_certo_sai_com_0(receita, tmp_path):
    a = _csv(tmp_path / "a.csv", [("Sul", "R$ 1,00")])
    b = _csv(tmp_path / "b.csv", [("Norte", "R$ 2,00")])
    result = runner.invoke(app, ["run", str(receita), str(a), str(b), "-o", str(tmp_path / "saida")])
    assert result.exit_code == 0, result.output
    assert sorted(p.name for p in (tmp_path / "saida").iterdir()) == ["a.csv", "b.csv"]


def test_lote_com_saida_que_e_arquivo(receita, tmp_path):
    a = _csv(tmp_path / "a.csv", [("Sul", "R$ 1,00")])
    b = _csv(tmp_path / "b.csv", [("Norte", "R$ 2,00")])
    arquivo = tmp_path / "x.csv"
    arquivo.write_text("x")
    result = runner.invoke(app, ["run", str(receita), str(a), str(b), "-o", str(arquivo)])
    assert result.exit_code == 2
    assert "precisa ser uma pasta" in result.stderr


def test_lote_padrao_sem_arquivo(receita, tmp_path):
    result = runner.invoke(app, ["run", str(receita), str(tmp_path / "*.xlsx"), "-o", str(tmp_path)])
    assert result.exit_code == 3
    assert "Nenhum arquivo casa" in result.stderr


def test_lote_arquivo_ilegivel_nao_para_os_outros(receita, tmp_path):
    a = _csv(tmp_path / "a.csv", [("Sul", "R$ 1,00")])
    ruim = tmp_path / "ruim.xlsx"
    ruim.write_bytes(b"isto nao e uma planilha")
    result = runner.invoke(app, ["run", str(receita), str(ruim), str(a), "-o", str(tmp_path / "saida")])
    assert result.exit_code == 3
    assert (tmp_path / "saida" / "a.csv").exists()
