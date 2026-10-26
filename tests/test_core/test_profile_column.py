from datetime import date

import openpyxl
import polars as pl
import pytest

from wolfz.core.ingest import IngestError
from wolfz.core.profile import apply_header, load_table, profile_column, profile_table, sample_rows


def test_perfil_de_coluna_numerica():
    s = pl.Series("valor", ["R$ 1.234,56", "R$ 10,00", None, "#N/D", "R$ 10,00", "(5,00)"])
    p = profile_column(s)
    assert p.name == "valor"
    assert p.inferred_type == "moeda"
    assert p.rows == 6
    assert p.nulls == 1
    assert p.null_markers == 1
    assert p.null_pct == pytest.approx(1 / 6)
    assert p.null_marker_pct == pytest.approx(1 / 6)
    assert p.distinct == 5  # None conta como um valor distinto, como no Polars
    assert p.top_shapes[0] == ("a$ 99,99", 2)
    assert p.examples == ["R$ 1.234,56", "R$ 10,00", "(5,00)"]  # sem repetição nem marcador
    assert (p.minimum, p.maximum) == (-5.0, 1234.56)


def test_perfil_de_coluna_data():
    p = profile_column(pl.Series("d", ["31/12/2024", "01/01/2020", "15/06/2022"]))
    assert p.inferred_type == "data"
    assert (p.minimum, p.maximum) == (date(2020, 1, 1), date(2024, 12, 31))
    assert p.to_dict()["minimum"] == "2020-01-01"


def test_perfil_sem_minimo_para_texto_e_codigo():
    assert profile_column(pl.Series("c", ["001", "002"])).minimum is None
    assert profile_column(pl.Series("t", ["Vitória", "Serra"])).maximum is None


def test_perfil_de_coluna_vazia():
    p = profile_column(pl.Series("v", [None, None], dtype=pl.String))
    assert p.inferred_type == "vazio"
    assert p.null_pct == 1.0
    assert p.examples == [] and p.top_shapes == []


def test_perfil_de_coluna_sem_linhas():
    p = profile_column(pl.Series("v", [], dtype=pl.String))
    assert (p.rows, p.null_pct, p.null_marker_pct) == (0, 0.0, 0.0)


def test_coluna_nao_texto_e_tratada_como_texto():
    p = profile_column(pl.Series("n", [1, 2, 3]))
    assert p.inferred_type == "numero"
    assert (p.minimum, p.maximum) == (1.0, 3.0)


def test_amostra_deterministica_e_em_ordem():
    df = pl.DataFrame({"i": list(range(100))})
    a = sample_rows(df, 10)
    assert a.height == 10
    assert a.equals(sample_rows(df, 10))
    assert a["i"].is_sorted()
    assert sample_rows(df, 1000) is df


def test_perfil_da_tabela_indica_amostra():
    df = pl.DataFrame({"a": [str(i) for i in range(50)], "b": ["x"] * 50})
    p = profile_table(df.lazy(), sample_size=20)
    assert (p.total_rows, p.sampled_rows, p.is_sample) == (50, 20, True)
    assert [c.name for c in p.columns] == ["a", "b"]
    assert all(c.rows == 20 for c in p.columns)
    d = p.to_dict()
    assert d["is_sample"] is True
    assert d["columns"][0]["top_shapes"][0]["count"] > 0
    assert not profile_table(df).is_sample


# ---------------------------------------------------------------------------
# Carregamento para inspeção
# ---------------------------------------------------------------------------


def test_apply_header_primeira_linha_nao_vazia():
    grid = pl.DataFrame(
        {
            "B": [None, "Nome", "Ana", None, "Bia"],
            "C": [None, None, "1", None, "2"],
            "D": [None, "Nome", "x", None, "y"],
        }
    )
    body, linha = apply_header(grid)
    assert linha == 2
    assert body.columns == ["Nome", "C", "Nome_2"]  # vazio vira a letra; repetido ganha sufixo
    assert body.rows() == [("Ana", "1", "x"), ("Bia", "2", "y")]  # linha vazia sai


def test_apply_header_linha_escolhida():
    grid = pl.DataFrame({"A": ["Título", "col", "1"]})
    body, linha = apply_header(grid, 2)
    assert (linha, body.columns, body.rows()) == (2, ["col"], [("1",)])
    with pytest.raises(ValueError, match="não existe"):
        apply_header(grid, 9)
    with pytest.raises(ValueError, match="vazia"):
        apply_header(pl.DataFrame())


def test_apply_header_grade_toda_vazia():
    body, linha = apply_header(pl.DataFrame({"A": [None, None]}, schema={"A": pl.String}))
    assert linha == 1 and body.height == 0


@pytest.fixture
def planilha(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dados"
    ws["B2"] = "Relatório"
    ws["B4"], ws["C4"] = "Município", "Valor"
    ws["B5"], ws["C5"] = "Vitória", "1.234,56"
    ws["B6"], ws["C6"] = "Serra", "10,00"
    wb.create_sheet("Outra")["A1"] = "x"
    path = tmp_path / "rel.xlsx"
    wb.save(path)
    return path


def test_load_table_planilha(planilha):
    t = load_table(planilha, header_row=4)
    assert (t.format, t.sheet, t.header_row) == ("xlsx", "Dados", 4)
    assert t.frame.columns == ["Município", "Valor"]
    assert [s.name for s in t.sheets] == ["Dados", "Outra"]
    assert load_table(planilha).header_row == 2  # sem indicação: a primeira não vazia
    assert load_table(planilha, sheet="Outra").frame.columns == ["x"]


def test_load_table_csv(tmp_path):
    path = tmp_path / "a.csv"
    path.write_text("a;b\n1;2\n", encoding="utf-8")
    t = load_table(path)
    assert (t.format, t.sheet, t.header_row, t.frame.columns) == ("csv", None, 1, ["a", "b"])
    assert load_table(path, header_row=1).frame.height == 1
    with pytest.raises(ValueError, match="abas"):
        load_table(path, sheet="x")
    with pytest.raises(ValueError, match="linha 1"):
        load_table(path, header_row=2)


def test_load_table_erro_de_leitura(tmp_path):
    with pytest.raises(IngestError):
        load_table(tmp_path / "nada.xlsx")


def test_minimo_e_maximo_de_coluna_ja_convertida():
    """Coluna Float relida como texto: '0.087' virava 87 (ponto de milhar), e
    o painel mostrava o CTR máximo como 87%."""
    p = profile_column(pl.Series("ctr", [0.087, 0.024, 0.09, None]))
    assert (p.minimum, p.maximum) == (0.024, 0.09)
    p = profile_column(pl.Series("peso", [1.234, 29.996]))
    assert (p.minimum, p.maximum) == (1.234, 29.996)
    p = profile_column(pl.Series("data", [date(2026, 1, 31), date(2025, 2, 1)]))
    assert (p.minimum, p.maximum) == (date(2025, 2, 1), date(2026, 1, 31))
