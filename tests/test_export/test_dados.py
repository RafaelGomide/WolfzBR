"""Exportação de dados (F1.I): CSV para o Excel brasileiro, Parquet, Excel
com tipos, build-then-swap e neutralização de fórmulas (F0.A9)."""

from datetime import date, datetime, time
from pathlib import Path

import polars as pl
import pytest
from openpyxl import load_workbook

from wolfz.export import dados
from wolfz.export.dados import (
    ExportError,
    OpcoesCsv,
    celulas_formula,
    escrever_csv,
    escrever_parquet,
    escrever_xlsx,
    escritor,
    gravar_juntos,
    neutralizar_formulas,
)

BOM = b"\xef\xbb\xbf"


@pytest.fixture
def tabela():
    return pl.DataFrame(
        {
            "nome": ["Ana", "=SOMA(A1:A9)", "+55 27 99999-0000", None],
            "valor": [1234.5, -0.25, None, float("nan")],
            "qtd": [1, 2, 3, None],
            "data": [date(2024, 1, 31), date(2023, 12, 1), None, date(2000, 2, 29)],
            "ok": [True, False, None, True],
        }
    )


# ---------------------------------------------------------------------------
# CSV (F1.I1)
# ---------------------------------------------------------------------------


def test_csv_tem_bom_e_virgula_por_padrao(tabela, tmp_path):
    escrever_csv(tabela, tmp_path / "a.csv")
    conteudo = (tmp_path / "a.csv").read_bytes()
    assert conteudo.startswith(BOM + b"nome,valor,qtd,data,ok\n")
    assert b"1234.5" in conteudo and b"2024-01-31" in conteudo


def test_csv_para_excel_brasileiro(tabela, tmp_path):
    escrever_csv(tabela, tmp_path / "a.csv", OpcoesCsv(separator=";", decimal_comma=True))
    linhas = (tmp_path / "a.csv").read_text(encoding="utf-8-sig").splitlines()
    assert linhas[0] == "nome;valor;qtd;data;ok"
    assert linhas[1] == "Ana;1234,5;1;2024-01-31;true"
    assert linhas[2].startswith("=SOMA(A1:A9);-0,25;2;")


def test_csv_virgula_decimal_com_separador_virgula_vai_entre_aspas(tmp_path):
    escrever_csv(pl.DataFrame({"v": [1.5]}), tmp_path / "a.csv", OpcoesCsv(decimal_comma=True))
    assert (tmp_path / "a.csv").read_text(encoding="utf-8-sig").splitlines()[1] == '"1,5"'


def test_csv_com_tabulacao(tmp_path):
    escrever_csv(pl.DataFrame({"a": ["x"], "b": [1]}), tmp_path / "a.tsv", OpcoesCsv(separator="\t"))
    assert (tmp_path / "a.tsv").read_text(encoding="utf-8-sig").splitlines() == ["a\tb", "x\t1"]


def test_csv_volta_igual_pelo_polars(tabela, tmp_path):
    escrever_csv(tabela, tmp_path / "a.csv", OpcoesCsv(separator=";"))
    lido = pl.read_csv(tmp_path / "a.csv", separator=";", try_parse_dates=True)
    assert lido.columns == tabela.columns  # o BOM não gruda no nome da primeira coluna
    assert lido["qtd"].to_list() == [1, 2, 3, None]


def test_separador_invalido():
    with pytest.raises(ValueError, match="separador"):
        OpcoesCsv(separator="|")  # type: ignore[arg-type]


# ---------------------------------------------------------------------------
# Parquet (F1.I2)
# ---------------------------------------------------------------------------


def test_parquet_preserva_tipos(tabela, tmp_path):
    escrever_parquet(tabela, tmp_path / "a.parquet")
    lido = pl.read_parquet(tmp_path / "a.parquet")
    assert lido.schema == tabela.schema
    assert lido.equals(tabela, null_equal=True)


# ---------------------------------------------------------------------------
# Excel (F1.I3)
# ---------------------------------------------------------------------------


def _celulas(path: Path, aba: str | None = None):
    wb = load_workbook(path)
    ws = wb[aba] if aba else wb.active
    return ws, [[c for c in linha] for linha in ws.iter_rows()]


def test_xlsx_com_tipos_corretos(tabela, tmp_path):
    escrever_xlsx(tabela, tmp_path / "a.xlsx", aba="vendas")
    ws, linhas = _celulas(tmp_path / "a.xlsx", "vendas")
    assert [c.value for c in linhas[0]] == ["nome", "valor", "qtd", "data", "ok"]
    assert all(c.font.bold for c in linhas[0])
    nome, valor, qtd, data, ok = linhas[1]
    assert nome.value == "Ana" and nome.data_type == "s"
    assert valor.value == 1234.5 and valor.data_type == "n"
    assert qtd.value == 1 and qtd.data_type == "n"
    assert data.value == datetime(2024, 1, 31) and data.is_date
    assert data.number_format == "DD/MM/YYYY"
    assert ok.value is True
    assert ws.freeze_panes == "A2"


def test_xlsx_texto_nunca_vira_formula(tabela, tmp_path):
    escrever_xlsx(tabela, tmp_path / "a.xlsx")
    _, linhas = _celulas(tmp_path / "a.xlsx")
    celula = linhas[2][0]
    assert celula.value == "=SOMA(A1:A9)" and celula.data_type == "s"


def test_xlsx_nulos_nan_e_vazios(tabela, tmp_path):
    escrever_xlsx(tabela, tmp_path / "a.xlsx")
    _, linhas = _celulas(tmp_path / "a.xlsx")
    assert [c.value for c in linhas[4]] == [None, None, None, datetime(2000, 2, 29), True]  # NaN vira vazio
    assert linhas[3][1].value is None


def test_xlsx_data_hora_e_hora(tmp_path):
    df = pl.DataFrame(
        {
            "quando": [datetime(2024, 5, 1, 13, 30)],
            "hora": [time(8, 15)],
            "lista": [[1, 2]],
        }
    )
    escrever_xlsx(df, tmp_path / "a.xlsx")
    _, linhas = _celulas(tmp_path / "a.xlsx")
    quando, hora, lista = linhas[1]
    assert quando.value == datetime(2024, 5, 1, 13, 30) and quando.number_format == "DD/MM/YYYY HH:MM:SS"
    assert hora.value == time(8, 15)
    assert lista.value == "[1, 2]"  # tipo sem equivalente no Excel vira texto


def test_xlsx_remove_caracteres_de_controle(tmp_path):
    escrever_xlsx(pl.DataFrame({"a": ["x\x01y\x1fz"]}), tmp_path / "a.xlsx")
    assert _celulas(tmp_path / "a.xlsx")[1][1][0].value == "xyz"


def test_xlsx_nome_da_aba_e_saneado(tmp_path):
    escrever_xlsx(pl.DataFrame({"a": [1]}), tmp_path / "a.xlsx", aba="Receita [2024]: jan/fev? " + "x" * 40)
    nome = load_workbook(tmp_path / "a.xlsx").sheetnames[0]
    assert len(nome) <= 31 and not set("[]:*?/\\") & set(nome)


def test_xlsx_acima_do_limite_de_linhas(monkeypatch, tmp_path):
    monkeypatch.setattr(dados, "MAX_LINHAS_XLSX", 3)
    with pytest.raises(ExportError, match="Excel aceita até 2"):
        escrever_xlsx(pl.DataFrame({"a": [1, 2, 3]}), tmp_path / "a.xlsx")
    assert not (tmp_path / "a.xlsx").exists()


def test_xlsx_texto_longo_demais(tmp_path):
    with pytest.raises(ExportError, match=r"coluna 'a'.*linha 2"):
        escrever_xlsx(pl.DataFrame({"a": ["x" * 40_000]}), tmp_path / "a.xlsx")


# ---------------------------------------------------------------------------
# Build-then-swap (F1.I4)
# ---------------------------------------------------------------------------


def test_gravar_juntos_grava_tudo(tmp_path):
    destinos = gravar_juntos(
        {
            tmp_path / "a.txt": lambda p: p.write_text("a"),
            tmp_path / "sub" / "b.txt": lambda p: p.write_text("b"),
        }
    )
    assert [p.read_text() for p in destinos] == ["a", "b"]
    assert not list(tmp_path.glob(".wolfz-export-*")) and not list((tmp_path / "sub").glob(".wolfz-*"))


def test_gravar_juntos_falha_nao_toca_nos_existentes(tmp_path):
    (tmp_path / "a.csv").write_text("antigo")

    def quebra(p: Path) -> None:
        p.write_text("metade")
        raise RuntimeError("disco cheio")

    with pytest.raises(RuntimeError, match="disco cheio"):
        gravar_juntos({tmp_path / "a.csv": lambda p: p.write_text("novo"), tmp_path / "b.csv": quebra})
    assert (tmp_path / "a.csv").read_text() == "antigo"  # o primeiro já estava pronto, mas não foi movido
    assert not (tmp_path / "b.csv").exists()
    assert sorted(p.name for p in tmp_path.iterdir()) == ["a.csv"]  # sem temporário esquecido


def test_xlsx_que_nao_cabe_nao_deixa_arquivo(monkeypatch, tmp_path):
    monkeypatch.setattr(dados, "MAX_LINHAS_XLSX", 2)
    df = pl.DataFrame({"a": [1, 2]})
    with pytest.raises(ExportError):
        gravar_juntos({tmp_path / "t.csv": escritor("csv", df), tmp_path / "t.xlsx": escritor("xlsx", df)})
    assert list(tmp_path.iterdir()) == []


# ---------------------------------------------------------------------------
# Fórmulas (F1.I5, F0.A9)
# ---------------------------------------------------------------------------


def test_conta_celulas_que_parecem_formula(tabela):
    assert celulas_formula(tabela) == 2  # "=SOMA…" e "+55…"; número negativo não é texto
    df = pl.DataFrame({"a": ["@x", "-1", "\tq", "\rq", "ok", "a=b", None]})
    assert celulas_formula(df) == 4
    assert celulas_formula(pl.DataFrame({"n": [1]})) == 0
    assert celulas_formula(pl.DataFrame({"a": []}, schema={"a": pl.String})) == 0


def test_neutralizar_prefixa_so_o_texto_perigoso(tabela):
    saida = neutralizar_formulas(tabela)
    assert saida["nome"].to_list() == ["Ana", "'=SOMA(A1:A9)", "'+55 27 99999-0000", None]
    assert saida["valor"].to_list()[:2] == [1234.5, -0.25]  # número fica como está
    assert celulas_formula(saida) == 0


def test_neutralizar_desligado_por_padrao(tabela, tmp_path):
    escritor("csv", tabela)(tmp_path / "a.csv")
    assert "\n=SOMA(A1:A9)," in (tmp_path / "a.csv").read_text(encoding="utf-8-sig")


def test_neutralizar_no_csv(tabela, tmp_path):
    escritor("csv", tabela, neutralizar=True)(tmp_path / "a.csv")
    texto = (tmp_path / "a.csv").read_text(encoding="utf-8-sig")
    assert "\n'=SOMA(A1:A9)," in texto and "\n'+55 27 99999-0000," in texto


def test_neutralizar_nao_muda_xlsx_nem_parquet(tabela, tmp_path):
    # No .xlsx o texto já é texto: o apóstrofo só sujaria o valor.
    escritor("xlsx", tabela, neutralizar=True)(tmp_path / "a.xlsx")
    assert _celulas(tmp_path / "a.xlsx")[1][2][0].value == "=SOMA(A1:A9)"
    escritor("parquet", tabela, neutralizar=True)(tmp_path / "a.parquet")
    assert pl.read_parquet(tmp_path / "a.parquet")["nome"][1] == "=SOMA(A1:A9)"


def test_formato_desconhecido():
    with pytest.raises(ValueError, match="formato"):
        escritor("docx", pl.DataFrame())
