"""Gera os arquivos sintéticos do corpus de testes (F1.B3 e F1.B4).

Cada arquivo imita a **estrutura** de uma fonte real (relatório de
prefeitura, tabela do SIDRA, exportação de ERP…) com valores inventados.
Nenhum dado de cliente ou de pessoa real entra no Git: CPFs e CNPJs são
gerados com dígito verificador válido, mas aleatórios.

Para cada arquivo, grava em ``tests/corpus/<nome>/``:

- ``original.<ext>``: a planilha suja;
- ``saida_esperada.parquet``: a tabela limpa, construída a partir dos dados
  que geraram o arquivo (é o gabarito, não passa pelos parsers do WolfzBR);
- ``receita_esperada.wolfz.json``: os passos que levam de um ao outro;
- ``FONTE.md``: origem, licença e categorias do corpus.

Arquivos com duas tabelas têm também ``saida_esperada_2.parquet`` e
``receita_esperada_2.wolfz.json``.

Receitas que usam operações ainda não implementadas são rascunho: quando a
operação entrar no catálogo, o teste do corpus passa a executá-la e mostra
se a receita (ou a operação) precisa de ajuste.

Determinístico: o mesmo nome gera sempre os mesmos dados. O ``.xls`` precisa
do pacote ``xlwt``, que não é dependência do projeto::

    python scripts/gerar_sinteticos.py              # todos
    python scripts/gerar_sinteticos.py anp_glp      # só um
    python scripts/gerar_sinteticos.py --listar

Para gerar os ``.xls``: ``pip install xlwt`` num ambiente à parte e rode com
``PYTHONPATH`` apontando para ele.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import random
import sys
import zipfile
import zlib
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl
import polars as pl
from openpyxl.styles import Font

RAIZ = Path(__file__).resolve().parents[1]
CORPUS = RAIZ / "tests" / "corpus"
DATA_FIXA = datetime(2026, 1, 1)
CRIADO_EM = "2026-10-16T00:00:00Z"

# --------------------------------------------------------------------------
# Dados de apoio (inventados ou de domínio público, como nomes de municípios)
# --------------------------------------------------------------------------

MUNICIPIOS_ES = [
    ("3205309", "Vitória"),
    ("3205200", "Vila Velha"),
    ("3205002", "Serra"),
    ("3201308", "Cariacica"),
    ("3201209", "Cachoeiro de Itapemirim"),
    ("3203205", "Linhares"),
    ("3204906", "São Mateus"),
    ("3201506", "Colatina"),
    ("3202405", "Guarapari"),
    ("3200607", "Aracruz"),
    ("3205101", "Viana"),
    ("3203908", "Nova Venécia"),
]
NOMES = ["Ana", "Bruno", "Carla", "Diego", "Eduarda", "Felipe", "Gabriela", "Heitor", "Isabela", "João"]
NOMES += ["Larissa", "Marcos", "Natália", "Otávio", "Patrícia", "Rafael", "Sabrina", "Tiago", "Vanessa"]
SOBRENOMES = ["Silva", "Souza", "Oliveira", "Santos", "Pereira", "Lima", "Carvalho", "Ferreira", "Almeida"]
SOBRENOMES += ["Rodrigues", "Nascimento", "Araújo", "Ribeiro", "Gomes", "Barbosa"]
EMPRESAS = ["Comercial", "Distribuidora", "Construtora", "Serviços", "Transportes", "Alimentos", "Tecnologia"]
EMPRESAS_2 = ["Capixaba", "Atlântico", "Vale do Rio Doce", "Litoral Norte", "Serrana", "Moxuara", "Itapoã"]


def _rng(nome: str) -> random.Random:
    return random.Random(zlib.crc32(nome.encode()))


def pessoa(rng: random.Random) -> str:
    return f"{rng.choice(NOMES)} {rng.choice(SOBRENOMES)} {rng.choice(SOBRENOMES)}"


def empresa(rng: random.Random) -> str:
    return f"{rng.choice(EMPRESAS)} {rng.choice(EMPRESAS_2)} Ltda"


def dinheiro(rng: random.Random, lo: float, hi: float) -> float:
    return round(rng.uniform(lo, hi), 2)


def br(v: float) -> str:
    """1234.5 -> '1.234,50'."""
    return f"{v:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")


def data_aleatoria(rng: random.Random, ano: int, ate_mes: int = 12) -> date:
    return date(ano, rng.randint(1, ate_mes), rng.randint(1, 28))


def _dv(digitos: str, pesos: list[int]) -> str:
    resto = sum(int(d) * p for d, p in zip(digitos, pesos, strict=True)) % 11
    return "0" if resto < 2 else str(11 - resto)


def cpf(rng: random.Random, zero_a_esquerda: bool = False) -> str:
    base = ("0" if zero_a_esquerda else str(rng.randint(1, 9))) + "".join(
        str(rng.randint(0, 9)) for _ in range(8)
    )
    d1 = _dv(base, list(range(10, 1, -1)))
    return base + d1 + _dv(base + d1, list(range(11, 1, -1)))


def cnpj(rng: random.Random, zeros_a_esquerda: int = 0) -> str:
    base = (
        "0" * zeros_a_esquerda + "".join(str(rng.randint(0, 9)) for _ in range(8 - zeros_a_esquerda)) + "0001"
    )
    pesos = [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]
    d1 = _dv(base, pesos)
    return base + d1 + _dv(base + d1, [6, *pesos])


def fmt_cpf_mascarado(c: str) -> str:
    """Como o Portal da Transparência publica (LGPD): ***.456.789-**."""
    return f"***.{c[3:6]}.{c[6:9]}-**"


def fmt_cnpj(c: str) -> str:
    return f"{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}"


def mojibake(texto: str) -> str:
    """UTF-8 lido como cp1252: 'ção' -> 'Ã§Ã£o'."""
    return texto.encode("utf-8").decode("cp1252")


# --------------------------------------------------------------------------
# Escrita dos formatos
# --------------------------------------------------------------------------


class Aba:
    """Grade de uma aba, 0-based."""

    def __init__(self, nome: str, oculta: bool = False):
        self.nome = nome
        self.oculta = oculta
        self.celulas: dict[tuple[int, int], object] = {}
        self.negrito: set[tuple[int, int]] = set()
        self.mesclas: list[tuple[int, int, int, int]] = []

    def linha(self, r: int, valores: list, c: int = 0, negrito: bool = False) -> Aba:
        for j, v in enumerate(valores):
            if v is None or v == "":
                continue
            self.celulas[(r, c + j)] = v
            if negrito:
                self.negrito.add((r, c + j))
        return self

    def mesclar(self, r1: int, c1: int, r2: int, c2: int) -> Aba:
        self.mesclas.append((r1, c1, r2, c2))
        return self


def escrever_xlsx(path: Path, abas: list[Aba]) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for aba in abas:
        ws = wb.create_sheet(aba.nome)
        for (r, c), v in sorted(aba.celulas.items()):
            cell = ws.cell(r + 1, c + 1, v)
            if (r, c) in aba.negrito:
                cell.font = Font(bold=True)
            if isinstance(v, date):
                cell.number_format = "DD/MM/YYYY"
        for r1, c1, r2, c2 in aba.mesclas:
            ws.merge_cells(start_row=r1 + 1, start_column=c1 + 1, end_row=r2 + 1, end_column=c2 + 1)
        if aba.oculta:
            ws.sheet_state = "hidden"
    wb.properties.created = wb.properties.modified = DATA_FIXA
    wb.save(path)


def escrever_xls(path: Path, abas: list[Aba]) -> None:
    try:
        import xlwt
    except ImportError as exc:
        raise SystemExit("Para gerar .xls instale xlwt num ambiente à parte (ver docstring).") from exc
    estilos = {
        (False, False): xlwt.XFStyle(),
        (True, False): xlwt.easyxf("font: bold on"),
        (False, True): xlwt.easyxf(num_format_str="DD/MM/YYYY"),
        (True, True): xlwt.easyxf("font: bold on", num_format_str="DD/MM/YYYY"),
    }
    wb = xlwt.Workbook(encoding="utf-8")
    for aba in abas:
        ws = wb.add_sheet(aba.nome)
        mesclas = {(m[0], m[1]): m for m in aba.mesclas}
        for (r, c), v in sorted(aba.celulas.items()):
            estilo = estilos[((r, c) in aba.negrito, isinstance(v, date))]
            if (r, c) in mesclas:
                r1, c1, r2, c2 = mesclas[(r, c)]
                ws.write_merge(r1, r2, c1, c2, v, estilo)
            else:
                ws.write(r, c, v, estilo)
        if aba.oculta:
            ws.visibility = 1
    wb.save(str(path))


_ODS_NS = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"'
)


def _ods_celula(v: object, attrs: str) -> str:
    if v is None:
        return f"<table:table-cell{attrs}/>"
    if isinstance(v, date):
        return (
            f'<table:table-cell office:value-type="date" office:date-value="{v:%Y-%m-%d}"{attrs}>'
            f"<text:p>{v:%d/%m/%Y}</text:p></table:table-cell>"
        )
    if isinstance(v, int | float):
        return (
            f'<table:table-cell office:value-type="float" office:value="{v}"{attrs}>'
            f"<text:p>{v}</text:p></table:table-cell>"
        )
    return (
        f'<table:table-cell office:value-type="string"{attrs}>'
        f"<text:p>{escape(str(v))}</text:p></table:table-cell>"
    )


def escrever_ods(path: Path, abas: list[Aba]) -> None:
    """ODS mínimo, com mesclagens (o ``calamine`` ainda não as lê) e bytes
    determinísticos."""
    tabelas = []
    for aba in abas:
        coberta = {
            (r, c)
            for r1, c1, r2, c2 in aba.mesclas
            for r in range(r1, r2 + 1)
            for c in range(c1, c2 + 1)
            if (r, c) != (r1, c1)
        }
        spans = {(m[0], m[1]): m for m in aba.mesclas}
        n_lin = max(r for r, _ in aba.celulas) + 1
        n_col = max(c for _, c in aba.celulas) + 1
        linhas = []
        for r in range(n_lin):
            cells = []
            for c in range(n_col):
                if (r, c) in coberta:
                    cells.append("<table:covered-table-cell/>")
                    continue
                attrs = ""
                if (r, c) in spans:
                    r1, c1, r2, c2 = spans[(r, c)]
                    attrs = (
                        f' table:number-columns-spanned="{c2 - c1 + 1}"'
                        f' table:number-rows-spanned="{r2 - r1 + 1}"'
                    )
                cells.append(_ods_celula(aba.celulas.get((r, c)), attrs))
            linhas.append(f"<table:table-row>{''.join(cells)}</table:table-row>")
        tabelas.append(f'<table:table table:name="{escape(aba.nome)}">{"".join(linhas)}</table:table>')
    content = (
        f'<?xml version="1.0" encoding="UTF-8"?><office:document-content {_ODS_NS} office:version="1.2">'
        f"<office:body><office:spreadsheet>{''.join(tabelas)}</office:spreadsheet></office:body>"
        "</office:document-content>"
    )
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">'
        '<manifest:file-entry manifest:full-path="/" '
        'manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/>'
        '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
        "</manifest:manifest>"
    )
    with zipfile.ZipFile(path, "w") as zf:
        for nome, dados, metodo in [
            ("mimetype", "application/vnd.oasis.opendocument.spreadsheet", zipfile.ZIP_STORED),
            ("META-INF/manifest.xml", manifest, zipfile.ZIP_DEFLATED),
            ("content.xml", content, zipfile.ZIP_DEFLATED),
        ]:
            info = zipfile.ZipInfo(nome, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = metodo
            info.create_system = 0  # o padrão depende do SO (0 no Windows, 3 no Linux)
            zf.writestr(info, dados)


def escrever_csv(
    path: Path, linhas: list[list], encoding: str = "utf-8", sep: str = ";", aspas: bool = False
):
    buf = io.StringIO()
    quoting = csv.QUOTE_ALL if aspas else csv.QUOTE_MINIMAL
    writer = csv.writer(buf, delimiter=sep, quoting=quoting, lineterminator="\r\n")
    for linha in linhas:
        writer.writerow(["" if v is None else v for v in linha])
    path.write_bytes(buf.getvalue().encode(encoding))


# --------------------------------------------------------------------------
# Modelo de um arquivo do corpus
# --------------------------------------------------------------------------


@dataclass
class Saida:
    tabela: pl.DataFrame
    passos: list[dict]
    aba: str | None = None


@dataclass
class Sintetico:
    nome: str
    categorias: tuple[str, ...]
    formato: str
    imita: str
    descricao: str
    escrever: Callable[[Path], None]
    saidas: list[Saida] = field(default_factory=list)


SINTETICOS: dict[str, Callable[[], Sintetico]] = {}


def sintetico(fn: Callable[[], Sintetico]) -> Callable[[], Sintetico]:
    SINTETICOS[fn.__name__] = fn
    return fn


def tabela(*colunas: tuple[str, pl.DataType, list]) -> pl.DataFrame:
    return pl.DataFrame({n: v for n, _, v in colunas}, schema={n: t for n, t, _ in colunas})


TXT, NUM, INT, DATA = pl.String, pl.Float64, pl.Int64, pl.Date


def receita_json(s: Sintetico, saida: Saida) -> str:
    source: dict = {"type": {"xlsx": "excel", "xls": "excel", "ods": "ods", "csv": "csv"}[s.formato]}
    if saida.aba:
        source["sheet"] = saida.aba
    receita = {
        "wolfz_version": "0.1",
        "name": s.nome.replace("_", " "),
        "created_at": CRIADO_EM,
        "source": source,
        "steps": [{"id": f"s{i}", **p} for i, p in enumerate(saida.passos, 1)],
    }
    return json.dumps(receita, ensure_ascii=False, indent=2, sort_keys=True) + "\n"


def fonte_md(s: Sintetico) -> str:
    return (
        "---\n"
        "origem: sintetico\n"
        "gerador: scripts/gerar_sinteticos.py\n"
        "licenca: MIT (dados inventados)\n"
        f"categorias: {', '.join(s.categorias)}\n"
        "---\n\n"
        f"# {s.nome}\n\n"
        f"{s.descricao}\n\n"
        f"Imita a estrutura de: {s.imita}. Os valores são inventados.\n"
    )


def gravar(s: Sintetico, destino: Path = CORPUS) -> Path:
    pasta = destino / s.nome
    pasta.mkdir(parents=True, exist_ok=True)
    for antigo in [
        *pasta.glob("original.*"),
        *pasta.glob("saida_esperada*"),
        *pasta.glob("receita_esperada*"),
    ]:
        antigo.unlink()
    s.escrever(pasta / f"original.{s.formato}")
    for i, saida in enumerate(s.saidas):
        sufixo = "" if i == 0 else f"_{i + 1}"
        saida.tabela.write_parquet(pasta / f"saida_esperada{sufixo}.parquet")
        (pasta / f"receita_esperada{sufixo}.wolfz.json").write_text(receita_json(s, saida), encoding="utf-8")
    (pasta / "FONTE.md").write_text(fonte_md(s), encoding="utf-8")
    return pasta


# --------------------------------------------------------------------------
# Relatórios com título, cabeçalho e rodapé
# --------------------------------------------------------------------------


@sintetico
def receita_municipal() -> Sintetico:
    rng = _rng("receita_municipal")
    receitas = [
        ("1.1.1.2.50.0.0", "Imposto sobre a Propriedade Predial e Territorial Urbana - IPTU"),
        ("1.1.1.3.03.1.0", "Imposto sobre a Renda Retido na Fonte - Trabalho"),
        ("1.1.1.4.51.1.0", "Imposto sobre Serviços de Qualquer Natureza - ISS"),
        ("1.1.1.2.53.0.0", "Imposto sobre Transmissão Inter Vivos - ITBI"),
        ("1.1.2.1.01.0.0", "Taxa de Fiscalização de Funcionamento"),
        ("1.1.2.2.01.0.0", "Taxa de Coleta de Lixo"),
        ("1.2.4.0.00.0.0", "Contribuição para Custeio da Iluminação Pública"),
        ("1.3.2.1.00.1.1", "Remuneração de Depósitos Bancários"),
        ("1.7.1.1.51.0.0", "Cota-Parte do Fundo de Participação dos Municípios"),
        ("1.7.2.1.50.0.0", "Cota-Parte do ICMS"),
        ("1.7.2.1.51.0.0", "Cota-Parte do IPVA"),
        ("1.9.9.0.99.0.0", "Outras Receitas"),
    ]
    previsto = [dinheiro(rng, 50_000, 9_000_000) for _ in receitas]
    arrecadado = [round(p * rng.uniform(0.7, 1.15), 2) for p in previsto]
    cab = ["Código", "Descrição da Receita", "Previsão Atualizada (R$)", "Arrecadado (R$)"]

    def escrever(path: Path) -> None:
        aba = Aba("Receita")
        aba.linha(0, ["PREFEITURA MUNICIPAL DE SANTA LUZIA DO NORTE"], negrito=True).mesclar(0, 0, 0, 3)
        aba.linha(1, ["Receita Orçamentária Arrecadada - Exercício 2025"], negrito=True).mesclar(1, 0, 1, 3)
        aba.linha(3, cab, negrito=True)
        for i, ((cod, desc), p, a) in enumerate(zip(receitas, previsto, arrecadado, strict=True)):
            aba.linha(4 + i, [cod, desc, p, a])
        aba.linha(17, ["Fonte: Sistema Integrado de Contabilidade Pública"])
        aba.linha(18, ["Emitido em 05/01/2026 às 10:32"])
        escrever_xlsx(path, [aba])

    saida = tabela(
        ("codigo", TXT, [c for c, _ in receitas]),
        ("descricao", TXT, [d for _, d in receitas]),
        ("previsao_atualizada", NUM, previsto),
        ("arrecadado", NUM, arrecadado),
    )
    passos = [
        {"op": "select_range", "start": "A4", "end_detect": "auto"},
        {"op": "header_rows", "n": 1},
        {"op": "rename", "mapping": dict(zip(cab, saida.columns, strict=True))},
        {"op": "cast", "types": {"previsao_atualizada": "float", "arrecadado": "float"}},
    ]
    return Sintetico(
        "receita_municipal",
        ("relatorio_rodape",),
        "xlsx",
        "relatório de receita arrecadada de portal de transparência municipal",
        "Título em duas linhas mescladas, cabeçalho em negrito na linha 4, "
        "rodapé com fonte e data de emissão.",
        escrever,
        [Saida(saida, passos, "Receita")],
    )


@sintetico
def despesa_por_orgao() -> Sintetico:
    rng = _rng("despesa_por_orgao")
    elementos = [
        "3.1.90.11 - Vencimentos e Vantagens Fixas",
        "3.3.90.30 - Material de Consumo",
        "3.3.90.39 - Outros Serviços de Terceiros - PJ",
        "4.4.90.51 - Obras e Instalações",
    ]
    orgaos = [
        "Secretaria Municipal de Saúde",
        "Secretaria Municipal de Educação",
        "Secretaria Municipal de Obras",
    ]
    linhas = []
    for orgao in orgaos:
        for elemento in rng.sample(elementos, 3):
            emp = dinheiro(rng, 10_000, 2_000_000)
            liq = round(emp * rng.uniform(0.5, 1), 2)
            pago = round(liq * rng.uniform(0.6, 1), 2)
            linhas.append((orgao, elemento, emp, liq, pago))
    cab = ["Órgão", "Elemento de Despesa", "Empenhado", "Liquidado", "Pago"]

    def soma(rows: list, j: int) -> float:
        return round(sum(r[j] for r in rows), 2)

    def escrever(path: Path) -> None:
        aba = Aba("Despesa")
        aba.linha(0, ["Despesa por Órgão e Elemento - 1º Quadrimestre de 2025"], negrito=True).mesclar(
            0, 0, 0, 4
        )
        aba.linha(2, cab, negrito=True)
        r = 3
        for orgao in orgaos:
            grupo = [x for x in linhas if x[0] == orgao]
            for x in grupo:
                aba.linha(r, list(x))
                r += 1
            aba.linha(
                r, [f"Subtotal {orgao}", None, soma(grupo, 2), soma(grupo, 3), soma(grupo, 4)], negrito=True
            )
            r += 1
        aba.linha(r, ["Total Geral", None, soma(linhas, 2), soma(linhas, 3), soma(linhas, 4)], negrito=True)
        aba.linha(r + 2, ["Fonte: Sistema de Execução Orçamentária e Financeira"])
        escrever_xlsx(path, [aba])

    saida = tabela(
        ("orgao", TXT, [x[0] for x in linhas]),
        ("elemento", TXT, [x[1] for x in linhas]),
        ("empenhado", NUM, [x[2] for x in linhas]),
        ("liquidado", NUM, [x[3] for x in linhas]),
        ("pago", NUM, [x[4] for x in linhas]),
    )
    passos = [
        {"op": "select_range", "start": "A3", "end_detect": "auto"},
        {"op": "header_rows", "n": 1},
        {"op": "drop_rows_where", "col": "Órgão", "match": "starts_with", "value": "Subtotal"},
        {"op": "drop_rows_where", "col": "Órgão", "match": "equals", "value": "Total Geral"},
        {"op": "rename", "mapping": dict(zip(cab, saida.columns, strict=True))},
        {"op": "cast", "types": {"empenhado": "float", "liquidado": "float", "pago": "float"}},
    ]
    return Sintetico(
        "despesa_por_orgao",
        ("relatorio_rodape", "subtotais"),
        "xlsx",
        "relatório de execução da despesa por órgão (Tesouro Transparente, portais municipais)",
        "Subtotal em negrito ao fim de cada órgão, total geral e nota de fonte "
        "depois de uma linha em branco.",
        escrever,
        [Saida(saida, passos, "Despesa")],
    )


@sintetico
def convenios_estaduais() -> Sintetico:
    rng = _rng("convenios_estaduais")
    objetos = [
        "Pavimentação de vias",
        "Construção de unidade de saúde",
        "Reforma de escola",
        "Aquisição de ônibus",
    ]
    linhas = []
    for _ in range(10):
        valor = dinheiro(rng, 100_000, 3_000_000)
        linhas.append(
            (
                f"{rng.randint(100, 999)}/2025",
                rng.choice(MUNICIPIOS_ES)[1],
                rng.choice(objetos),
                valor,
                round(valor * rng.choice([0.01, 0.02, 0.05]), 2),
                data_aleatoria(rng, 2025),
            )
        )
    cab = ["Nº Convênio", "Município", "Objeto", "Valor (R$)", "Contrapartida (R$)", "Data de assinatura"]

    def escrever(path: Path) -> None:
        rows: list[list] = [
            ["GOVERNO DO ESTADO - SECRETARIA DE ECONOMIA E PLANEJAMENTO"],
            ["Convênios celebrados com municípios - 2025"],
            [],
            cab,
        ]
        rows += [[n, m, o, br(v), br(c), f"{d:%d/%m/%Y}"] for n, m, o, v, c, d in linhas]
        total_v, total_c = sum(x[3] for x in linhas), sum(x[4] for x in linhas)
        rows += [["Total", "", "", br(total_v), br(total_c), ""], [], ["Fonte: SIGEFES"]]
        escrever_csv(path, rows, encoding="cp1252")

    saida = tabela(
        ("numero_convenio", TXT, [x[0] for x in linhas]),
        ("municipio", TXT, [x[1] for x in linhas]),
        ("objeto", TXT, [x[2] for x in linhas]),
        ("valor", NUM, [x[3] for x in linhas]),
        ("contrapartida", NUM, [x[4] for x in linhas]),
        ("data_assinatura", DATA, [x[5] for x in linhas]),
    )
    passos = [
        {"op": "select_range", "start": "A4", "end_detect": "auto"},
        {"op": "header_rows", "n": 1},
        {"op": "drop_rows_where", "col": "Nº Convênio", "match": "equals", "value": "Total"},
        {"op": "rename", "mapping": dict(zip(cab, saida.columns, strict=True))},
        {"op": "parse_brl", "cols": ["valor", "contrapartida"]},
        {"op": "parse_date_br", "cols": ["data_assinatura"]},
    ]
    return Sintetico(
        "convenios_estaduais",
        ("relatorio_rodape", "csv_cp1252"),
        "csv",
        "relatório de convênios exportado em CSV por sistema estadual",
        "CSV cp1252 com ';', duas linhas de título antes do cabeçalho, linha Total e fonte no fim.",
        escrever,
        [Saida(saida, passos)],
    )


# --------------------------------------------------------------------------
# Cabeçalho multinível e sinais do IBGE
# --------------------------------------------------------------------------


def _sinal_ibge(rng: random.Random, lo: int, hi: int) -> tuple[object, float | None]:
    """Valor como o SIDRA publica e o valor limpo esperado."""
    sorte = rng.random()
    if sorte < 0.1:
        return "-", 0.0
    if sorte < 0.15:
        return "..", None
    if sorte < 0.2:
        return "...", None
    if sorte < 0.25:
        return "X", None
    v = rng.randint(lo, hi)
    return v, float(v)


@sintetico
def sidra_pam_area() -> Sintetico:
    rng = _rng("sidra_pam_area")
    municipios = [f"{m} (ES)" for _, m in MUNICIPIOS_ES[:10]]
    cols = [
        "area_plantada_2022",
        "area_plantada_2023",
        "quantidade_produzida_2022",
        "quantidade_produzida_2023",
    ]
    brutos, limpos = {c: [] for c in cols}, {c: [] for c in cols}
    for _ in municipios:
        for c in cols:
            b, v = _sinal_ibge(rng, 10, 50_000)
            brutos[c].append(b)
            limpos[c].append(v)
    area = "Área plantada ou destinada à colheita (Hectares)"
    qtd = "Quantidade produzida (Toneladas)"

    def escrever(path: Path) -> None:
        aba = Aba("Tabela 5457")
        aba.linha(
            0, ["Tabela 5457 - Área plantada e quantidade produzida das lavouras temporárias e permanentes"]
        )
        aba.linha(1, ["Município", area, None, qtd]).mesclar(1, 0, 2, 0).mesclar(1, 1, 1, 2).mesclar(
            1, 3, 1, 4
        )
        aba.linha(2, ["2022", "2023", "2022", "2023"], c=1)
        for i, m in enumerate(municipios):
            aba.linha(3 + i, [m, *(brutos[c][i] for c in cols)])
        r = 3 + len(municipios)
        aba.linha(r, ["Fonte: IBGE - Produção Agrícola Municipal"])
        aba.linha(r + 1, ["Notas"])
        aba.linha(r + 2, ["1 - Os dados do último ano divulgado são RESULTADOS PRELIMINARES."])
        aba.linha(r + 3, ["Legenda"])
        aba.linha(r + 4, ["-  - Dado numérico igual a zero não resultante de arredondamento"])
        aba.linha(r + 5, ["..  - Não se aplica dado numérico"])
        aba.linha(r + 6, ["... - Dado numérico não disponível"])
        aba.linha(r + 7, ["X  - Dado numérico omitido a fim de evitar a individualização da informação"])
        escrever_xlsx(path, [aba])

    saida = tabela(("municipio", TXT, municipios), *((c, NUM, limpos[c]) for c in cols))
    passos = [
        {"op": "select_range", "start": "A2", "end_detect": "auto"},
        {"op": "header_rows", "n": 2, "joiner": " "},
        {
            "op": "rename",
            "mapping": {
                "Município": "municipio",
                f"{area} 2022": cols[0],
                f"{area} 2023": cols[1],
                f"{qtd} 2022": cols[2],
                f"{qtd} 2023": cols[3],
            },
        },
        {"op": "parse_ibge", "cols": cols},
    ]
    return Sintetico(
        "sidra_pam_area",
        ("cabecalho_multinivel", "sinais_ibge"),
        "xlsx",
        "tabela do SIDRA baixada em xlsx (PAM, tabela 5457)",
        "Cabeçalho em duas linhas (variável mesclada sobre os anos), sinais '-', '..', '...' e 'X', "
        "notas e legenda logo abaixo dos dados, sem linha em branco.",
        escrever,
        [Saida(saida, passos, "Tabela 5457")],
    )


@sintetico
def boletim_saude_multinivel() -> Sintetico:
    rng = _rng("boletim_saude_multinivel")
    regioes = ["Central", "Metropolitana", "Norte", "Sul"]
    meses = ["Janeiro", "Fevereiro", "Março"]
    cols = [f"{m.lower().replace('ç', 'c')}_{s}" for m in meses for s in ("masculino", "feminino")]
    dados = {c: [rng.randint(0, 900) for _ in regioes] for c in cols}

    def escrever(path: Path) -> None:
        aba = Aba("Dengue")
        aba.linha(0, ["Casos notificados de dengue por sexo - Espírito Santo, 2024"], negrito=True)
        aba.mesclar(0, 0, 0, 6)
        aba.linha(2, ["Região de Saúde"], negrito=True).mesclar(2, 0, 3, 0)
        for j, m in enumerate(meses):
            aba.linha(2, [m], c=1 + 2 * j, negrito=True).mesclar(2, 1 + 2 * j, 2, 2 + 2 * j)
            aba.linha(3, ["Masculino", "Feminino"], c=1 + 2 * j, negrito=True)
        for i, reg in enumerate(regioes):
            aba.linha(4 + i, [reg, *(dados[c][i] for c in cols)])
        aba.linha(8, ["Total", *(sum(dados[c]) for c in cols)], negrito=True)
        aba.linha(10, ["Fonte: SESA/ES - dados sujeitos a revisão."])
        escrever_xlsx(path, [aba])

    saida = tabela(("regiao_de_saude", TXT, regioes), *((c, INT, dados[c]) for c in cols))
    mapping = {"Região de Saúde": "regiao_de_saude"}
    mapping |= {
        f"{m} {s.capitalize()}": f"{m.lower().replace('ç', 'c')}_{s}"
        for m in meses
        for s in ("masculino", "feminino")
    }
    passos = [
        {"op": "select_range", "start": "A3", "end_detect": "auto"},
        {"op": "header_rows", "n": 2, "joiner": " "},
        {"op": "drop_rows_where", "col": "Região de Saúde", "match": "equals", "value": "Total"},
        {"op": "rename", "mapping": mapping},
        {"op": "cast", "types": dict.fromkeys(cols, "int")},
    ]
    return Sintetico(
        "boletim_saude_multinivel",
        ("cabecalho_multinivel",),
        "xlsx",
        "boletim epidemiológico de secretaria estadual de saúde",
        "Mês mesclado sobre Masculino/Feminino, 'Região de Saúde' mesclada na vertical, total em negrito.",
        escrever,
        [Saida(saida, passos, "Dengue")],
    )


@sintetico
def sidra_ppm_rebanho() -> Sintetico:
    rng = _rng("sidra_ppm_rebanho")
    municipios = [f"{m} (ES)" for _, m in MUNICIPIOS_ES[2:12]]
    cab = ["Município", "Bovino", "Bubalino", "Equino", "Suíno - total", "Caprino", "Galináceos - total"]
    cols = ["municipio", "bovino", "bubalino", "equino", "suino_total", "caprino", "galinaceos_total"]
    brutos, limpos = {c: [] for c in cols[1:]}, {c: [] for c in cols[1:]}
    for _ in municipios:
        for c in cols[1:]:
            b, v = _sinal_ibge(rng, 5, 200_000)
            brutos[c].append(b)
            limpos[c].append(v)

    def escrever(path: Path) -> None:
        aba = Aba("Tabela 3939")
        aba.linha(0, ["Tabela 3939 - Efetivo dos rebanhos, por tipo de rebanho"])
        aba.linha(1, ["Variável - Efetivo dos rebanhos (Cabeças)"])
        aba.linha(2, ["Ano - 2023"])
        aba.linha(3, cab)
        for i, m in enumerate(municipios):
            aba.linha(4 + i, [m, *(brutos[c][i] for c in cols[1:])])
        r = 4 + len(municipios)
        aba.linha(r, ["Fonte: IBGE - Pesquisa da Pecuária Municipal"])
        aba.linha(r + 1, ["Legenda"])
        aba.linha(r + 2, ["-  - Dado numérico igual a zero não resultante de arredondamento"])
        aba.linha(r + 3, ["X  - Dado numérico omitido a fim de evitar a individualização da informação"])
        escrever_xlsx(path, [aba])

    saida = tabela(("municipio", TXT, municipios), *((c, NUM, limpos[c]) for c in cols[1:]))
    passos = [
        {"op": "select_range", "start": "A4", "end_detect": "auto"},
        {"op": "header_rows", "n": 1},
        {"op": "rename", "mapping": dict(zip(cab, cols, strict=True))},
        {"op": "parse_ibge", "cols": cols[1:]},
    ]
    return Sintetico(
        "sidra_ppm_rebanho",
        ("sinais_ibge",),
        "xlsx",
        "tabela do SIDRA baixada em xlsx (PPM, tabela 3939)",
        "Três linhas de título (tabela, variável, ano), sinais do IBGE nos valores, fonte e legenda no fim.",
        escrever,
        [Saida(saida, passos, "Tabela 3939")],
    )


@sintetico
def censo_populacao_sinais() -> Sintetico:
    rng = _rng("censo_populacao_sinais")
    linhas = []
    for cod, m in MUNICIPIOS_ES:
        homens, mulheres = rng.randint(5_000, 200_000), rng.randint(5_000, 200_000)
        if rng.random() < 0.15:
            linhas.append((cod, f"{m} (ES)", "...", "...", "...", None, None, None))
        else:
            linhas.append(
                (
                    cod,
                    f"{m} (ES)",
                    homens + mulheres,
                    homens,
                    mulheres,
                    *map(float, (homens + mulheres, homens, mulheres)),
                )
            )
    cab = ["Cód.", "Município", "Total", "Homens", "Mulheres"]

    def escrever(path: Path) -> None:
        rows: list[list] = [
            ["Tabela 9514 - População residente, por sexo"],
            ["Variável - População residente (Pessoas)"],
            ["Ano - 2022"],
            cab,
        ]
        rows += [list(x[:5]) for x in linhas]
        rows += [["Fonte: IBGE - Censo Demográfico"]]
        escrever_csv(path, rows, encoding="utf-8", aspas=True)

    saida = tabela(
        ("cod", TXT, [x[0] for x in linhas]),
        ("municipio", TXT, [x[1] for x in linhas]),
        ("total", NUM, [x[5] for x in linhas]),
        ("homens", NUM, [x[6] for x in linhas]),
        ("mulheres", NUM, [x[7] for x in linhas]),
    )
    passos = [
        {"op": "select_range", "start": "A4", "end_detect": "auto"},
        {"op": "header_rows", "n": 1},
        {"op": "rename", "mapping": dict(zip(cab, saida.columns, strict=True))},
        {"op": "parse_ibge", "cols": ["total", "homens", "mulheres"]},
    ]
    return Sintetico(
        "censo_populacao_sinais",
        ("sinais_ibge",),
        "csv",
        "tabela do SIDRA baixada em CSV (Censo 2022, tabela 9514)",
        "CSV UTF-8 com aspas em todos os campos, três linhas de título, '...' para dado indisponível.",
        escrever,
        [Saida(saida, passos)],
    )


# --------------------------------------------------------------------------
# Subtotais e mesclagem vertical
# --------------------------------------------------------------------------


@sintetico
def execucao_funcao_subfuncao() -> Sintetico:
    rng = _rng("execucao_funcao_subfuncao")
    funcoes = [
        (
            "10 - Saúde",
            [
                "301 - Atenção Básica",
                "302 - Assistência Hospitalar e Ambulatorial",
                "305 - Vigilância Epidemiológica",
            ],
        ),
        ("12 - Educação", ["361 - Ensino Fundamental", "365 - Educação Infantil"]),
        ("04 - Administração", ["122 - Administração Geral", "126 - Tecnologia da Informação"]),
        ("15 - Urbanismo", ["451 - Infraestrutura Urbana", "452 - Serviços Urbanos"]),
    ]
    itens = []
    for funcao, subs in funcoes:
        for sub in subs:
            dot = dinheiro(rng, 200_000, 20_000_000)
            emp = round(dot * rng.uniform(0.4, 0.95), 2)
            itens.append((funcao, sub, dot, emp, round(emp * rng.uniform(0.5, 1), 2)))
    cab = ["Função / Subfunção", "Dotação Atualizada", "Empenhado", "Liquidado"]

    def soma(rows: list, j: int) -> float:
        return round(sum(r[j] for r in rows), 2)

    def escrever(path: Path) -> None:
        aba = Aba("Execução")
        aba.linha(0, ["Execução Orçamentária por Função e Subfunção - 2025"], negrito=True).mesclar(
            0, 0, 0, 3
        )
        aba.linha(2, cab, negrito=True)
        r = 3
        for funcao, _ in funcoes:
            grupo = [x for x in itens if x[0] == funcao]
            aba.linha(r, [funcao, soma(grupo, 2), soma(grupo, 3), soma(grupo, 4)], negrito=True)
            for x in grupo:
                aba.linha(r + 1, [f"    {x[1]}", x[2], x[3], x[4]])
                r += 1
            r += 1
        aba.linha(r, ["TOTAL", soma(itens, 2), soma(itens, 3), soma(itens, 4)], negrito=True)
        aba.linha(r + 2, ["Fonte: Sistema de Contabilidade Pública - posição em 31/12/2025"])
        escrever_xlsx(path, [aba])

    saida = tabela(
        ("subfuncao", TXT, [x[1] for x in itens]),
        ("dotacao_atualizada", NUM, [x[2] for x in itens]),
        ("empenhado", NUM, [x[3] for x in itens]),
        ("liquidado", NUM, [x[4] for x in itens]),
    )
    passos = [
        {"op": "select_range", "start": "A3", "end_detect": "auto"},
        {"op": "header_rows", "n": 1},
        {
            "op": "drop_rows_where",
            "col": "Função / Subfunção",
            "match": "in",
            "value": [f for f, _ in funcoes] + ["TOTAL"],
        },
        {"op": "rename", "mapping": dict(zip(cab, saida.columns, strict=True))},
        {"op": "trim", "cols": ["subfuncao"]},
        {"op": "cast", "types": {"dotacao_atualizada": "float", "empenhado": "float", "liquidado": "float"}},
    ]
    return Sintetico(
        "execucao_funcao_subfuncao",
        ("subtotais",),
        "xlsx",
        "execução orçamentária por função e subfunção (RREO, portais de transparência)",
        "Linha de função em negrito com o subtotal antes das subfunções "
        "(recuadas com espaços) e TOTAL no fim.",
        escrever,
        [Saida(saida, passos, "Execução")],
    )


@sintetico
def vendas_regiao_mesclada() -> Sintetico:
    rng = _rng("vendas_regiao_mesclada")
    estrutura = {"Sudeste": ["ES", "MG", "RJ"], "Sul": ["PR", "SC"], "Nordeste": ["BA"]}
    linhas = []
    for regiao, ufs in estrutura.items():
        for uf in ufs:
            for _ in range(rng.randint(2, 3)):
                linhas.append((regiao, uf, pessoa(rng), dinheiro(rng, 5_000, 300_000)))
    cab = ["Região", "UF", "Vendedor", "Vendas (R$)"]

    def escrever(path: Path) -> None:
        aba = Aba("Vendas")
        aba.linha(0, cab, negrito=True)
        for i, x in enumerate(linhas):
            r = 1 + i
            primeira_regiao = i == 0 or linhas[i - 1][0] != x[0]
            primeira_uf = i == 0 or linhas[i - 1][1] != x[1]
            aba.linha(r, [x[0] if primeira_regiao else None, x[1] if primeira_uf else None, x[2], x[3]])
            if primeira_regiao:
                fim = r + sum(1 for y in linhas[i:] if y[0] == x[0]) - 1
                aba.mesclar(r, 0, fim, 0)
            if primeira_uf:
                fim = r + sum(1 for y in linhas[i:] if y[1] == x[1]) - 1
                if fim > r:
                    aba.mesclar(r, 1, fim, 1)
        escrever_xlsx(path, [aba])

    saida = tabela(
        ("regiao", TXT, [x[0] for x in linhas]),
        ("uf", TXT, [x[1] for x in linhas]),
        ("vendedor", TXT, [x[2] for x in linhas]),
        ("vendas", NUM, [x[3] for x in linhas]),
    )
    passos = [
        {"op": "header_rows", "n": 1},
        {"op": "fill_down", "cols": ["Região", "UF"]},
        {"op": "rename", "mapping": dict(zip(cab, saida.columns, strict=True))},
        {"op": "cast", "types": {"vendas": "float"}},
    ]
    return Sintetico(
        "vendas_regiao_mesclada",
        ("mesclagem_vertical",),
        "xlsx",
        "relatório gerencial de vendas por região",
        "Região e UF mescladas na vertical: só a primeira linha de cada grupo tem valor.",
        escrever,
        [Saida(saida, passos, "Vendas")],
    )


@sintetico
def estoque_categoria_mesclada() -> Sintetico:
    rng = _rng("estoque_categoria_mesclada")
    categorias = {
        "Limpeza": ["Detergente 500 ml", "Água sanitária 1 l", "Sabão em pó 1 kg"],
        "Escritório": ["Papel A4 resma", "Caneta azul cx 50", "Grampeador"],
        "Copa": ["Café 500 g", "Açúcar 1 kg", "Copo descartável 200 ml"],
    }
    linhas = [
        (cat, prod, rng.randint(0, 500), dinheiro(rng, 2, 80))
        for cat, prods in categorias.items()
        for prod in prods
    ]
    cab = ["Categoria", "Produto", "Quantidade", "Custo unitário (R$)"]

    def escrever(path: Path) -> None:
        aba = Aba("Estoque")
        aba.linha(0, ["Posição de estoque em 30/06/2025"], negrito=True).mesclar(0, 0, 0, 3)
        aba.linha(2, cab, negrito=True)
        r = 3
        for cat, prods in categorias.items():
            aba.mesclar(r, 0, r + len(prods) - 1, 0)
            for i, x in enumerate(x for x in linhas if x[0] == cat):
                aba.linha(r + i, [cat if i == 0 else None, x[1], x[2], x[3]])
            r += len(prods)
        escrever_xlsx(path, [aba])

    saida = tabela(
        ("categoria", TXT, [x[0] for x in linhas]),
        ("produto", TXT, [x[1] for x in linhas]),
        ("quantidade", INT, [x[2] for x in linhas]),
        ("custo_unitario", NUM, [x[3] for x in linhas]),
    )
    passos = [
        {"op": "select_range", "start": "A3", "end_detect": "auto"},
        {"op": "header_rows", "n": 1},
        {"op": "fill_down", "cols": ["Categoria"]},
        {"op": "rename", "mapping": dict(zip(cab, saida.columns, strict=True))},
        {"op": "cast", "types": {"quantidade": "int", "custo_unitario": "float"}},
    ]
    return Sintetico(
        "estoque_categoria_mesclada",
        ("mesclagem_vertical",),
        "xlsx",
        "planilha de controle de estoque de almoxarifado",
        "Título mesclado, categoria mesclada na vertical sobre os produtos.",
        escrever,
        [Saida(saida, passos, "Estoque")],
    )


@sintetico
def prefeitura_folha_xls() -> Sintetico:
    rng = _rng("prefeitura_folha_xls")
    secretarias = {
        "Administração": ["Auxiliar Administrativo", "Assistente Administrativo", "Contador"],
        "Educação": ["Professor MaPA", "Pedagogo", "Auxiliar de Creche"],
        "Saúde": ["Enfermeiro", "Técnico de Enfermagem", "Agente Comunitário de Saúde"],
    }
    linhas = [
        (sec, cargo, (n := rng.randint(1, 120)), round(n * dinheiro(rng, 1_500, 9_000), 2))
        for sec, cargos in secretarias.items()
        for cargo in cargos
    ]
    cab = ["Secretaria", "Cargo", "Servidores", "Total bruto (R$)"]

    def escrever(path: Path) -> None:
        aba = Aba("Folha")
        aba.linha(0, ["Folha de pagamento por cargo - junho/2025"], negrito=True).mesclar(0, 0, 0, 3)
        aba.linha(1, cab, negrito=True)
        r = 2
        for sec, cargos in secretarias.items():
            aba.mesclar(r, 0, r + len(cargos) - 1, 0)
            for i, x in enumerate(x for x in linhas if x[0] == sec):
                aba.linha(r + i, [sec if i == 0 else None, x[1], x[2], x[3]])
            r += len(cargos)
        escrever_xls(path, [aba])

    saida = tabela(
        ("secretaria", TXT, [x[0] for x in linhas]),
        ("cargo", TXT, [x[1] for x in linhas]),
        ("servidores", INT, [x[2] for x in linhas]),
        ("total_bruto", NUM, [x[3] for x in linhas]),
    )
    passos = [
        {"op": "select_range", "start": "A2", "end_detect": "auto"},
        {"op": "header_rows", "n": 1},
        {"op": "fill_down", "cols": ["Secretaria"]},
        {"op": "rename", "mapping": dict(zip(cab, saida.columns, strict=True))},
        {"op": "cast", "types": {"servidores": "int", "total_bruto": "float"}},
    ]
    return Sintetico(
        "prefeitura_folha_xls",
        ("xls_antigo", "mesclagem_vertical"),
        "xls",
        "planilha .xls antiga de folha de pagamento de prefeitura",
        "Formato binário .xls (BIFF8), secretaria mesclada na vertical.",
        escrever,
        [Saida(saida, passos, "Folha")],
    )


@sintetico
def tabnet_obitos() -> Sintetico:
    rng = _rng("tabnet_obitos")
    anos = ["2019", "2020", "2021", "2022", "2023"]
    linhas = []
    for cod, m in MUNICIPIOS_ES[:10]:
        valores = [rng.choice([0, 0, *range(1, 400)]) for _ in anos]
        linhas.append((f"{cod[:6]} {m}", valores))
    cab = ["Município", *anos, "Total"]

    def escrever(path: Path) -> None:
        aba = Aba("Óbitos")
        aba.linha(0, ["Óbitos p/Residênc por Ano do Óbito segundo Município"])
        aba.linha(1, ["UF: Espírito Santo"])
        aba.linha(2, ["Período: 2019-2023"])
        aba.linha(4, cab)
        for i, (m, vs) in enumerate(linhas):
            aba.linha(5 + i, [m, *(v if v else "-" for v in vs), sum(vs)])
        r = 5 + len(linhas)
        totais = [sum(vs[j] for _, vs in linhas) for j in range(len(anos))]
        aba.linha(r, ["Total", *totais, sum(totais)])
        aba.linha(r + 2, ["Fonte: MS/SVS/CGIAE - Sistema de Informações sobre Mortalidade - SIM"])
        aba.linha(r + 3, ["Notas:"])
        aba.linha(r + 4, ["1. Em 2023, os dados são preliminares."])
        escrever_xls(path, [aba])

    cols = [f"ano_{a}" for a in anos]
    saida = tabela(
        ("municipio", TXT, [m for m, _ in linhas]),
        *((c, INT, [vs[j] for _, vs in linhas]) for j, c in enumerate(cols)),
    )
    passos = [
        {"op": "select_range", "start": "A5", "end_detect": "auto"},
        {"op": "header_rows", "n": 1},
        {"op": "drop_rows_where", "col": "Município", "match": "equals", "value": "Total"},
        {"op": "drop_cols", "cols": ["Total"]},
        {"op": "rename", "mapping": {"Município": "municipio", **dict(zip(anos, cols, strict=True))}},
        {"op": "parse_ibge", "cols": cols},
        {"op": "cast", "types": dict.fromkeys(cols, "int")},
    ]
    return Sintetico(
        "tabnet_obitos",
        ("xls_antigo", "relatorio_rodape"),
        "xls",
        "tabela do DATASUS TabNet (SIM) salva como .xls",
        "Formato .xls, três linhas de título, '-' para zero, linha e coluna Total, fonte e notas no fim.",
        escrever,
        [Saida(saida, passos, "Óbitos")],
    )


# --------------------------------------------------------------------------
# CSV: encoding e BOM
# --------------------------------------------------------------------------


@sintetico
def ceis_sancoes() -> Sintetico:
    rng = _rng("ceis_sancoes")
    linhas = []
    for _ in range(12):
        inicio = data_aleatoria(rng, rng.randint(2018, 2024))
        final = (
            None if rng.random() < 0.3 else date(inicio.year + rng.randint(1, 5), inicio.month, inicio.day)
        )
        if rng.random() < 0.4:
            c = cpf(rng)
            linhas.append(("F", fmt_cpf_mascarado(c), f"***{c[3:9]}**", pessoa(rng).upper(), inicio, final))
        else:
            c = cnpj(rng)
            linhas.append(("J", fmt_cnpj(c), c, empresa(rng).upper(), inicio, final))
    codigos = [str(rng.randint(100_000, 999_999)) for _ in linhas]
    ufs = [rng.choice(["ES", "MG", "RJ", "BA", "SP"]) for _ in linhas]
    cab = [
        "CADASTRO",
        "CÓDIGO DA SANÇÃO",
        "TIPO DE PESSOA",
        "CPF OU CNPJ DO SANCIONADO",
        "NOME DO SANCIONADO",
        "DATA INÍCIO SANÇÃO",
        "DATA FINAL SANÇÃO",
        "UF ÓRGÃO SANCIONADOR",
    ]

    def escrever(path: Path) -> None:
        rows = [cab]
        for (tipo, doc, _, nome, ini, fim), cod, uf in zip(linhas, codigos, ufs, strict=True):
            rows.append(
                [
                    "CEIS",
                    cod,
                    tipo,
                    doc,
                    nome,
                    f"{ini:%d/%m/%Y}",
                    f"{fim:%d/%m/%Y}" if fim else "Sem informação",
                    uf,
                ]
            )
        escrever_csv(path, rows, encoding="cp1252", aspas=True)

    saida = tabela(
        ("cadastro", TXT, ["CEIS"] * len(linhas)),
        ("codigo_sancao", TXT, codigos),
        ("tipo_pessoa", TXT, [x[0] for x in linhas]),
        ("cpf_cnpj", TXT, [x[2] for x in linhas]),
        ("nome", TXT, [x[3] for x in linhas]),
        ("data_inicio", DATA, [x[4] for x in linhas]),
        ("data_final", DATA, [x[5] for x in linhas]),
        ("uf", TXT, ufs),
    )
    passos = [
        {"op": "rename", "mapping": dict(zip(cab, saida.columns, strict=True))},
        {"op": "clean_cpf_cnpj", "col": "cpf_cnpj", "person_type_col": "tipo_pessoa"},
        {"op": "parse_date_br", "cols": ["data_inicio", "data_final"]},
    ]
    return Sintetico(
        "ceis_sancoes",
        ("csv_cp1252",),
        "csv",
        "CSV do Cadastro de Empresas Inidôneas e Suspensas (CEIS), Portal da Transparência",
        "CSV cp1252 com ';' e aspas em tudo, CPF mascarado (LGPD), CNPJ formatado "
        "e 'Sem informação' como nulo.",
        escrever,
        [Saida(saida, passos)],
    )


@sintetico
def servidores_remuneracao() -> Sintetico:
    rng = _rng("servidores_remuneracao")
    linhas = []
    for _ in range(15):
        basica = dinheiro(rng, 2_000, 35_000)
        natalina = dinheiro(rng, 1_000, 15_000) if rng.random() < 0.2 else 0.0
        verbas = None if rng.random() < 0.25 else dinheiro(rng, 0, 3_000)
        linhas.append(
            (
                str(rng.randint(1_000_000, 9_999_999)),
                fmt_cpf_mascarado(cpf(rng)),
                pessoa(rng).upper(),
                basica,
                natalina,
                verbas,
            )
        )
    cab = [
        "ANO",
        "MES",
        "Id_SERVIDOR_PORTAL",
        "CPF",
        "NOME",
        "REMUNERAÇÃO BÁSICA BRUTA (R$)",
        "GRATIFICAÇÃO NATALINA (R$)",
        "TOTAL DE VERBAS INDENIZATÓRIAS (R$)",
    ]

    def escrever(path: Path) -> None:
        rows = [cab]
        for ident, c, nome, b, n, v in linhas:
            rows.append(["2025", "06", ident, c, nome, br(b), br(n), "Não se aplica" if v is None else br(v)])
        escrever_csv(path, rows, encoding="cp1252", aspas=True)

    cols = [
        "ano",
        "mes",
        "id_servidor",
        "cpf",
        "nome",
        "remuneracao_basica",
        "gratificacao_natalina",
        "verbas_indenizatorias",
    ]
    saida = tabela(
        ("ano", TXT, ["2025"] * len(linhas)),
        ("mes", TXT, ["06"] * len(linhas)),
        ("id_servidor", TXT, [x[0] for x in linhas]),
        ("cpf", TXT, [x[1] for x in linhas]),
        ("nome", TXT, [x[2] for x in linhas]),
        ("remuneracao_basica", NUM, [x[3] for x in linhas]),
        ("gratificacao_natalina", NUM, [x[4] for x in linhas]),
        ("verbas_indenizatorias", NUM, [x[5] for x in linhas]),
    )
    passos = [
        {"op": "rename", "mapping": dict(zip(cab, cols, strict=True))},
        {"op": "parse_brl", "cols": cols[5:]},
    ]
    return Sintetico(
        "servidores_remuneracao",
        ("csv_cp1252",),
        "csv",
        "CSV de remuneração de servidores do Portal da Transparência",
        "CSV cp1252 com ';' e aspas em tudo, valores em formato brasileiro "
        "e 'Não se aplica' como nulo textual.",
        escrever,
        [Saida(saida, passos)],
    )


def _anp_linhas(rng: random.Random, produtos: list[tuple[str, float, float, str]], n: int) -> list[tuple]:
    bandeiras = ["BRANCA", "IPIRANGA", "VIBRA ENERGIA", "RAIZEN", "ALESAT"]
    linhas = []
    for _ in range(n):
        _, municipio = rng.choice(MUNICIPIOS_ES)
        produto, lo, hi, unidade = rng.choice(produtos)
        linhas.append(
            (
                "SE",
                "ES",
                municipio.upper().replace("Ó", "O").replace("Ã", "A").replace("Ê", "E").replace("É", "E"),
                f"POSTO {rng.choice(EMPRESAS_2).upper()} LTDA",
                cnpj(rng),
                produto,
                data_aleatoria(rng, 2025, 6),
                round(rng.uniform(lo, hi), 2),
                unidade,
                rng.choice(bandeiras),
            )
        )
    return linhas


@sintetico
def anp_precos() -> Sintetico:
    rng = _rng("anp_precos")
    produtos = [
        ("GASOLINA", 5.9, 6.9, "R$ / litro"),
        ("ETANOL", 3.9, 4.9, "R$ / litro"),
        ("DIESEL S10", 5.8, 6.6, "R$ / litro"),
    ]
    linhas = _anp_linhas(rng, produtos, 15)
    cab = [
        "Regiao - Sigla",
        "Estado - Sigla",
        "Municipio",
        "Revenda",
        "CNPJ da Revenda",
        "Nome da Rua",
        "Numero Rua",
        "Complemento",
        "Bairro",
        "Cep",
        "Produto",
        "Data da Coleta",
        "Valor de Venda",
        "Valor de Compra",
        "Unidade de Medida",
        "Bandeira",
    ]

    def escrever(path: Path) -> None:
        rows = [cab]
        for reg, uf, mun, rev, c, prod, d, v, un, band in linhas:
            rows.append(
                [reg, uf, mun, rev, fmt_cnpj(c), "AVENIDA BRASIL", str(rng.randint(1, 3000)), "", "CENTRO",
                 "29050-000", prod, f"{d:%d/%m/%Y}", f"{v:.2f}".replace(".", ","), "", un, band]
            )  # fmt: skip
        escrever_csv(path, rows, encoding="utf-8-sig")

    saida = tabela(
        ("regiao", TXT, [x[0] for x in linhas]),
        ("uf", TXT, [x[1] for x in linhas]),
        ("municipio", TXT, [x[2] for x in linhas]),
        ("revenda", TXT, [x[3] for x in linhas]),
        ("cnpj", TXT, [x[4] for x in linhas]),
        ("produto", TXT, [x[5] for x in linhas]),
        ("data_coleta", DATA, [x[6] for x in linhas]),
        ("valor_venda", NUM, [x[7] for x in linhas]),
        ("valor_compra", NUM, [None] * len(linhas)),
        ("unidade", TXT, [x[8] for x in linhas]),
        ("bandeira", TXT, [x[9] for x in linhas]),
    )
    passos = [
        {"op": "drop_cols", "cols": ["Nome da Rua", "Numero Rua", "Complemento", "Bairro", "Cep"]},
        {
            "op": "rename",
            "mapping": {
                "Regiao - Sigla": "regiao",
                "Estado - Sigla": "uf",
                "Municipio": "municipio",
                "Revenda": "revenda",
                "CNPJ da Revenda": "cnpj",
                "Produto": "produto",
                "Data da Coleta": "data_coleta",
                "Valor de Venda": "valor_venda",
                "Valor de Compra": "valor_compra",
                "Unidade de Medida": "unidade",
                "Bandeira": "bandeira",
            },
        },
        {"op": "clean_cnpj", "cols": ["cnpj"]},
        {"op": "parse_date_br", "cols": ["data_coleta"]},
        {"op": "parse_decimal_br", "cols": ["valor_venda", "valor_compra"]},
    ]
    return Sintetico(
        "anp_precos",
        ("csv_utf8_bom",),
        "csv",
        "série histórica de preços de combustíveis da ANP",
        "CSV UTF-8 com BOM (o BOM gruda no nome da primeira coluna se lido como UTF-8 puro), "
        "decimal com vírgula e 'Valor de Compra' sempre vazio.",
        escrever,
        [Saida(saida, passos)],
    )


@sintetico
def anp_glp() -> Sintetico:
    rng = _rng("anp_glp")
    linhas = _anp_linhas(rng, [("GLP", 95.0, 125.0, "R$ / 13 kg")], 10)
    cab = [
        "Regiao - Sigla",
        "Estado - Sigla",
        "Municipio",
        "Revenda",
        "Produto",
        "Data da Coleta",
        "Valor de Venda",
        "Unidade de Medida",
        "Bandeira",
    ]
    cols = [
        "regiao",
        "uf",
        "municipio",
        "revenda",
        "produto",
        "data_coleta",
        "valor_venda",
        "unidade",
        "bandeira",
    ]

    def escrever(path: Path) -> None:
        rows = [cab]
        for reg, uf, mun, rev, _, prod, d, v, un, band in linhas:
            rows.append([reg, uf, mun, rev, prod, f"{d:%d/%m/%Y}", f"{v:.2f}".replace(".", ","), un, band])
        escrever_csv(path, rows, encoding="utf-8-sig")

    saida = tabela(
        ("regiao", TXT, [x[0] for x in linhas]),
        ("uf", TXT, [x[1] for x in linhas]),
        ("municipio", TXT, [x[2] for x in linhas]),
        ("revenda", TXT, [x[3] for x in linhas]),
        ("produto", TXT, [x[5] for x in linhas]),
        ("data_coleta", TXT, [f"{x[6]:%d/%m/%Y}" for x in linhas]),
        ("valor_venda", NUM, [x[7] for x in linhas]),
        ("unidade", TXT, [x[8] for x in linhas]),
        ("bandeira", TXT, [x[9] for x in linhas]),
    )
    passos = [
        {"op": "rename", "mapping": dict(zip(cab, cols, strict=True))},
        {"op": "parse_brl", "cols": ["valor_venda"]},
    ]
    return Sintetico(
        "anp_glp",
        ("csv_utf8_bom",),
        "csv",
        "levantamento de preços de GLP da ANP",
        "CSV UTF-8 com BOM e ';'. A receita usa só operações já implementadas; a data fica como texto.",
        escrever,
        [Saida(saida, passos)],
    )


# --------------------------------------------------------------------------
# Datas em 8 dígitos e zeros perdidos
# --------------------------------------------------------------------------


@sintetico
def cnpj_estabelecimentos() -> Sintetico:
    rng = _rng("cnpj_estabelecimentos")
    municipios_receita = ["5705", "5703", "5699", "5625", "5623", "5663"]
    linhas = []
    for _ in range(12):
        c = cnpj(rng, zeros_a_esquerda=rng.choice([0, 0, 1, 2]))
        situacao = rng.choice(["02", "02", "04", "08"])
        inicio = data_aleatoria(rng, rng.randint(1990, 2023))
        data_sit = (
            None if rng.random() < 0.2 else data_aleatoria(rng, rng.randint(max(inicio.year, 2005), 2025))
        )
        cnae = rng.choice(["4711302", "5611201", "4930202", "8599604"])
        cep = f"29{rng.randint(0, 999999):06d}"
        mun = rng.choice(municipios_receita)
        linhas.append(
            (
                c[:8],
                c[8:12],
                c[12:],
                "1",
                empresa(rng).upper(),
                situacao,
                data_sit,
                "00",
                inicio,
                cnae,
                cep,
                "ES",
                mun,
            )
        )
    nomes = [
        "cnpj_basico",
        "cnpj_ordem",
        "cnpj_dv",
        "identificador_matriz_filial",
        "nome_fantasia",
        "situacao_cadastral",
        "data_situacao_cadastral",
        "motivo_situacao_cadastral",
        "data_inicio_atividade",
        "cnae_fiscal_principal",
        "cep",
        "uf",
        "municipio",
    ]

    def escrever(path: Path) -> None:
        rows = []
        for x in linhas:
            sit = f"{x[6]:%Y%m%d}" if x[6] else rng.choice(["0", ""])
            rows.append([*x[:6], sit, x[7], f"{x[8]:%Y%m%d}", *x[9:]])
        escrever_csv(path, rows, encoding="cp1252", aspas=True)

    tipos = [TXT] * len(nomes)
    tipos[6] = tipos[8] = DATA
    saida = tabela(
        *((n, t, [x[i] for x in linhas]) for i, (n, t) in enumerate(zip(nomes, tipos, strict=True)))
    )
    passos = [
        {"op": "header_rows", "n": 0, "names": nomes},
        {
            "op": "parse_date_br",
            "cols": ["data_situacao_cadastral", "data_inicio_atividade"],
            "eight_digits": "ymd",
        },
    ]
    return Sintetico(
        "cnpj_estabelecimentos",
        ("datas_8_digitos", "csv_cp1252"),
        "csv",
        "arquivo de Estabelecimentos dos Dados Abertos do CNPJ (Receita Federal)",
        "CSV sem cabeçalho, cp1252, ';' e aspas em tudo; datas AAAAMMDD, com '0' ou vazio como data ausente.",
        escrever,
        [Saida(saida, passos)],
    )


@sintetico
def sinasc_nascimentos() -> Sintetico:
    rng = _rng("sinasc_nascimentos")
    linhas = []
    for _ in range(15):
        nasc = data_aleatoria(rng, 2021)
        cadastro = date(2021, min(nasc.month + 1, 12), rng.randint(1, 28))
        linhas.append(
            (
                rng.choice(MUNICIPIOS_ES)[0][:6],
                nasc,
                rng.choice(["1", "2"]),
                rng.randint(1_800, 4_500),
                cadastro,
            )
        )
    cab = ["CODMUNNASC", "DTNASC", "SEXO", "PESO", "DTCADASTRO"]

    def escrever(path: Path) -> None:
        rows = [cab]
        for mun, nasc, sexo, peso, cad in linhas:
            # Passou pelo Excel: DDMMAAAA com dia < 10 perde o zero à esquerda.
            dt = f"{nasc:%d%m%Y}".lstrip("0")
            rows.append([mun, dt, sexo, str(peso), f"{cad:%d%m%Y}"])
        escrever_csv(path, rows, sep=",")

    saida = tabela(
        ("codmunnasc", TXT, [x[0] for x in linhas]),
        ("dtnasc", DATA, [x[1] for x in linhas]),
        ("sexo", TXT, [x[2] for x in linhas]),
        ("peso", INT, [x[3] for x in linhas]),
        ("dtcadastro", DATA, [x[4] for x in linhas]),
    )
    passos = [
        {"op": "rename", "mapping": dict(zip(cab, saida.columns, strict=True))},
        {"op": "parse_date_br", "cols": ["dtnasc", "dtcadastro"], "eight_digits": "dmy"},
        {"op": "cast", "types": {"peso": "int"}},
    ]
    return Sintetico(
        "sinasc_nascimentos",
        ("datas_8_digitos", "zeros_perdidos"),
        "csv",
        "microdados do SINASC (DATASUS) exportados para CSV",
        "Datas DDMMAAAA; na DTNASC o zero à esquerda do dia se perdeu (7 dígitos) em parte das linhas.",
        escrever,
        [Saida(saida, passos)],
    )


@sintetico
def clientes_cpf_cep() -> Sintetico:
    rng = _rng("clientes_cpf_cep")
    ceps_sp = ["01310100", "04538133", "05402000", "02012000", "08010000"]
    linhas = []
    for _ in range(12):
        c = cpf(rng, zero_a_esquerda=rng.random() < 0.5)
        cep = rng.choice(ceps_sp) if rng.random() < 0.5 else f"29{rng.randint(0, 999999):06d}"
        linhas.append(
            (
                pessoa(rng),
                c,
                cep,
                "São Paulo" if cep[0] == "0" else "Vitória",
                "SP" if cep[0] == "0" else "ES",
            )
        )
    cab = ["Nome", "CPF", "CEP", "Cidade", "UF"]

    def escrever(path: Path) -> None:
        aba = Aba("Clientes")
        aba.linha(0, cab, negrito=True)
        for i, (nome, c, cep, cidade, uf) in enumerate(linhas):
            # Parte digitada como número (perde o zero), parte como texto formatado.
            valor_cpf = int(c) if i % 3 else f"{c[:3]}.{c[3:6]}.{c[6:9]}-{c[9:]}"
            aba.linha(1 + i, [nome, valor_cpf, int(cep), cidade, uf])
        escrever_xlsx(path, [aba])

    saida = tabela(
        ("nome", TXT, [x[0] for x in linhas]),
        ("cpf", TXT, [x[1] for x in linhas]),
        ("cep", TXT, [x[2] for x in linhas]),
        ("cidade", TXT, [x[3] for x in linhas]),
        ("uf", TXT, [x[4] for x in linhas]),
    )
    passos = [
        {"op": "header_rows", "n": 1},
        {"op": "rename", "mapping": dict(zip(cab, saida.columns, strict=True))},
        {"op": "clean_cpf", "cols": ["cpf"]},
        {"op": "clean_cep", "cols": ["cep"]},
    ]
    return Sintetico(
        "clientes_cpf_cep",
        ("zeros_perdidos",),
        "xlsx",
        "cadastro de clientes mantido no Excel",
        "CPF e CEP digitados como número perdem o zero à esquerda; parte dos CPFs está como texto formatado.",
        escrever,
        [Saida(saida, passos, "Clientes")],
    )


@sintetico
def fornecedores_cnpj() -> Sintetico:
    rng = _rng("fornecedores_cnpj")
    linhas = []
    for _ in range(12):
        c = cnpj(rng, zeros_a_esquerda=rng.choice([0, 1, 2, 3]))
        cep = rng.choice(["01310100", "04538133", f"29{rng.randint(0, 999999):06d}"])
        linhas.append((empresa(rng), c, cep, rng.choice(MUNICIPIOS_ES)[1], "ES"))
    cab = ["Razão Social", "CNPJ", "CEP", "Município", "UF"]

    def escrever(path: Path) -> None:
        rows = [cab] + [[r, str(int(c)), str(int(cep)), m, uf] for r, c, cep, m, uf in linhas]
        escrever_csv(path, rows, encoding="cp1252")

    saida = tabela(
        ("razao_social", TXT, [x[0] for x in linhas]),
        ("cnpj", TXT, [x[1] for x in linhas]),
        ("cep", TXT, [x[2] for x in linhas]),
        ("municipio", TXT, [x[3] for x in linhas]),
        ("uf", TXT, [x[4] for x in linhas]),
    )
    passos = [
        {"op": "rename", "mapping": dict(zip(cab, saida.columns, strict=True))},
        {"op": "clean_cnpj", "cols": ["cnpj"]},
        {"op": "clean_cep", "cols": ["cep"]},
    ]
    return Sintetico(
        "fornecedores_cnpj",
        ("zeros_perdidos", "csv_cp1252"),
        "csv",
        "cadastro de fornecedores salvo como CSV pelo Excel",
        "CSV cp1252 com ';'; CNPJ e CEP sem os zeros à esquerda.",
        escrever,
        [Saida(saida, passos)],
    )


# --------------------------------------------------------------------------
# Meses lado a lado
# --------------------------------------------------------------------------


@sintetico
def bcb_serie_mensal() -> Sintetico:
    rng = _rng("bcb_serie_mensal")
    indicadores = [
        ("IPCA - variação mensal", "%", -0.5, 1.2),
        ("IGP-M - variação mensal", "%", -1.0, 1.5),
        ("Selic - taxa média", "% a.a.", 10.0, 11.5),
        ("Câmbio - dólar (venda, média)", "R$/US$", 4.8, 5.6),
    ]
    meses = ["jan", "fev", "mar", "abr", "mai", "jun"]
    valores = {nome: [round(rng.uniform(lo, hi), 2) for _ in meses] for nome, _, lo, hi in indicadores}
    cab = ["Indicador", "Unidade", *(f"{m}/2024" for m in meses)]

    def escrever(path: Path) -> None:
        aba = Aba("Indicadores")
        aba.linha(0, ["Indicadores econômicos selecionados - 2024"], negrito=True)
        aba.linha(2, cab, negrito=True)
        for i, (nome, unidade, _, _) in enumerate(indicadores):
            aba.linha(3 + i, [nome, unidade, *valores[nome]])
        aba.linha(8, ["Fonte: Banco Central do Brasil - Sistema Gerenciador de Séries Temporais (SGS)"])
        escrever_xlsx(path, [aba])

    saida = tabela(
        ("indicador", TXT, [n for n, *_ in indicadores for _ in meses]),
        ("unidade", TXT, [u for _, u, *_ in indicadores for _ in meses]),
        ("mes", DATA, [date(2024, j + 1, 1) for _ in indicadores for j in range(len(meses))]),
        ("valor", NUM, [v for n, *_ in indicadores for v in valores[n]]),
    )
    passos = [
        {"op": "select_range", "start": "A3", "end_detect": "auto"},
        {"op": "header_rows", "n": 1},
        {"op": "rename", "mapping": {"Indicador": "indicador", "Unidade": "unidade"}},
        {"op": "unpivot", "index": ["indicador", "unidade"], "variable_name": "mes", "value_name": "valor"},
        {"op": "parse_month_year", "cols": ["mes"]},
        {"op": "cast", "types": {"valor": "float"}},
    ]
    return Sintetico(
        "bcb_serie_mensal",
        ("meses_lado_a_lado",),
        "xlsx",
        "planilha de séries mensais do Banco Central",
        "Um indicador por linha e um mês por coluna ('jan/2024'); "
        "a saída é longa, uma linha por indicador e mês.",
        escrever,
        [Saida(saida, passos, "Indicadores")],
    )


@sintetico
def orcamento_gerencial_meses() -> Sintetico:
    rng = _rng("orcamento_gerencial_meses")
    linhas = [
        ("Administrativo", "Aluguel"),
        ("Administrativo", "Energia elétrica"),
        ("Comercial", "Comissões"),
        ("Comercial", "Publicidade"),
        ("Operações", "Manutenção"),
        ("Operações", "Combustível"),
    ]
    meses = ["Jan", "Fev", "Mar", "Abr"]
    valores = [[dinheiro(rng, 1_000, 60_000) for _ in meses] for _ in linhas]
    cab = ["Centro de Custo", "Conta", *meses, "Total"]

    def escrever(path: Path) -> None:
        aba = Aba("Orçamento")
        aba.linha(0, ["Orçamento 2025 - Despesas por Centro de Custo"], negrito=True)
        aba.linha(2, cab, negrito=True)
        for i, ((cc, conta), vs) in enumerate(zip(linhas, valores, strict=True)):
            aba.linha(3 + i, [cc, conta, *vs, round(sum(vs), 2)])
        totais = [round(sum(vs[j] for vs in valores), 2) for j in range(len(meses))]
        aba.linha(3 + len(linhas), ["Total", None, *totais, round(sum(totais), 2)], negrito=True)
        escrever_xlsx(path, [aba])

    saida = tabela(
        ("centro_custo", TXT, [cc for cc, _ in linhas for _ in meses]),
        ("conta", TXT, [c for _, c in linhas for _ in meses]),
        ("mes", TXT, meses * len(linhas)),
        ("valor", NUM, [v for vs in valores for v in vs]),
    )
    passos = [
        {"op": "select_range", "start": "A3", "end_detect": "auto"},
        {"op": "header_rows", "n": 1},
        {"op": "drop_cols", "cols": ["Total"]},
        {"op": "drop_rows_where", "col": "Centro de Custo", "match": "equals", "value": "Total"},
        {"op": "rename", "mapping": {"Centro de Custo": "centro_custo", "Conta": "conta"}},
        {"op": "unpivot", "index": ["centro_custo", "conta"], "variable_name": "mes", "value_name": "valor"},
        {"op": "cast", "types": {"valor": "float"}},
    ]
    return Sintetico(
        "orcamento_gerencial_meses",
        ("meses_lado_a_lado", "subtotais"),
        "xlsx",
        "orçamento gerencial mensal de empresa",
        "Meses 'Jan'..'Abr' lado a lado, coluna Total à direita e linha Total no fim.",
        escrever,
        [Saida(saida, passos, "Orçamento")],
    )


# --------------------------------------------------------------------------
# Exportação de ERP
# --------------------------------------------------------------------------


@sintetico
def erp_sap_lancamentos() -> Sintetico:
    rng = _rng("erp_sap_lancamentos")
    textos = ["Pagamento fornecedor", "Recebimento cliente", "Estorno de lançamento", "Adiantamento"]
    linhas = []
    for _ in range(14):
        lanc = data_aleatoria(rng, 2024)
        venc = None if rng.random() < 0.3 else date(lanc.year, min(lanc.month + 1, 12), lanc.day)
        montante = dinheiro(rng, 10, 90_000) * rng.choice([1, -1])
        linhas.append(
            (
                "1000",
                f"{rng.randint(100_000_000, 199_999_999):010d}",
                rng.choice(["KR", "DZ", "SA"]),
                lanc,
                montante,
                "BRL",
                venc,
                rng.choice(textos),
            )
        )
    cab = [
        "Empresa",
        "Nº documento",
        "Tipo doc.",
        "Data lançamento",
        "Montante em MI",
        "Moeda",
        "Vencimento líquido",
        "Texto",
    ]

    def escrever(path: Path) -> None:
        rows = [cab]
        for emp, doc, tipo, lanc, mont, moeda, venc, texto in linhas:
            valor = br(abs(mont)) + ("-" if mont < 0 else "")
            rows.append(
                [
                    emp,
                    doc,
                    tipo,
                    f"{lanc:%d.%m.%Y}",
                    valor,
                    moeda,
                    f"{venc:%d.%m.%Y}" if venc else "31.12.9999",
                    texto,
                ]
            )
        escrever_csv(path, rows, encoding="cp1252")

    cols = ["empresa", "documento", "tipo_doc", "data_lancamento", "montante", "moeda", "vencimento", "texto"]
    tipos = [TXT, TXT, TXT, DATA, NUM, TXT, DATA, TXT]
    saida = tabela(
        *((c, t, [x[i] for x in linhas]) for i, (c, t) in enumerate(zip(cols, tipos, strict=True)))
    )
    passos = [
        {"op": "rename", "mapping": dict(zip(cab, cols, strict=True))},
        {"op": "replace", "cols": ["vencimento"], "old": "31.12.9999", "new": None},
        {"op": "parse_date_br", "cols": ["data_lancamento", "vencimento"]},
        {"op": "parse_brl", "cols": ["montante"]},
    ]
    return Sintetico(
        "erp_sap_lancamentos",
        ("exportacao_erp", "csv_cp1252"),
        "csv",
        "exportação de partidas de um ERP SAP (transação FBL1N/FBL5N) para CSV",
        "Sinal de menos no fim ('1.234,56-'), datas 'DD.MM.AAAA' e '31.12.9999' como vencimento em aberto.",
        escrever,
        [Saida(saida, passos)],
    )


@sintetico
def erp_totvs_razao() -> Sintetico:
    rng = _rng("erp_totvs_razao")
    historicos = [
        "Venda de mercadorias",
        "Pagamento a fornecedor",
        "Depósito bancário",
        "Pagamento de tarifa",
    ]
    saldo = 5_000.0
    linhas = []
    for dia in sorted(rng.sample(range(2, 31), 12)):
        valor = dinheiro(rng, 50, 8_000)
        debito = rng.random() < 0.5
        saldo = round(saldo + (valor if debito else -valor), 2)
        linhas.append(
            (
                date(2025, 1, dia),
                f"{rng.randint(1, 999):06d}",
                rng.choice(historicos),
                valor if debito else None,
                None if debito else valor,
                saldo,
            )
        )

    def saldo_txt(v: float) -> str:
        return f"{br(abs(v))} {'D' if v >= 0 else 'C'}"

    cab = ["Data", "Lote", "Histórico", "Débito", "Crédito", "Saldo"]

    def escrever(path: Path) -> None:
        aba = Aba("Razão")
        aba.linha(0, ["TOTVS - Razão Contábil"], negrito=True)
        aba.linha(1, ["Empresa: 01 - Matriz"])
        aba.linha(2, ["Conta: 1.1.01.001 - Caixa Geral"])
        aba.linha(3, ["Período: 01/01/2025 a 31/01/2025"])
        aba.linha(5, cab, negrito=True)
        aba.linha(6, [None, None, "SALDO ANTERIOR", None, None, saldo_txt(5_000.0)])
        for i, (d, lote, hist, deb, cred, s) in enumerate(linhas):
            aba.linha(
                7 + i,
                [
                    f"{d:%d/%m/%Y}",
                    lote,
                    hist,
                    br(deb) if deb else None,
                    br(cred) if cred else None,
                    saldo_txt(s),
                ],
            )
        total_d = round(sum(x[3] or 0 for x in linhas), 2)
        total_c = round(sum(x[4] or 0 for x in linhas), 2)
        aba.linha(
            7 + len(linhas),
            [None, None, "TOTAL DA CONTA", br(total_d), br(total_c), saldo_txt(saldo)],
            negrito=True,
        )
        escrever_xlsx(path, [aba])

    saida = tabela(
        ("data", DATA, [x[0] for x in linhas]),
        ("lote", TXT, [x[1] for x in linhas]),
        ("historico", TXT, [x[2] for x in linhas]),
        ("debito", NUM, [x[3] for x in linhas]),
        ("credito", NUM, [x[4] for x in linhas]),
        # debit_negative: saldo devedor (D) negativo, credor (C) positivo.
        ("saldo", NUM, [-x[5] for x in linhas]),
    )
    passos = [
        {"op": "select_range", "start": "A6", "end_detect": "auto"},
        {"op": "header_rows", "n": 1},
        {
            "op": "drop_rows_where",
            "col": "Histórico",
            "match": "in",
            "value": ["SALDO ANTERIOR", "TOTAL DA CONTA"],
        },
        {"op": "rename", "mapping": dict(zip(cab, saida.columns, strict=True))},
        {"op": "parse_date_br", "cols": ["data"]},
        {"op": "parse_brl", "cols": ["debito", "credito"]},
        {"op": "parse_brl", "cols": ["saldo"], "debit_negative": True},
    ]
    return Sintetico(
        "erp_totvs_razao",
        ("exportacao_erp",),
        "xlsx",
        "relatório Razão Contábil exportado de ERP TOTVS para Excel",
        "Tudo como texto, quatro linhas de cabeçalho do relatório, saldo com sufixo D/C, "
        "saldo anterior e total da conta.",
        escrever,
        [Saida(saida, passos, "Razão")],
    )


# --------------------------------------------------------------------------
# ODS
# --------------------------------------------------------------------------


@sintetico
def transparencia_contratos_ods() -> Sintetico:
    rng = _rng("transparencia_contratos_ods")
    objetos = [
        "Fornecimento de merenda escolar",
        "Locação de veículos",
        "Manutenção predial",
        "Serviços de limpeza",
    ]
    linhas = [
        (
            f"{i + 1:03d}/2025",
            empresa(rng),
            cnpj(rng),
            rng.choice(objetos),
            dinheiro(rng, 20_000, 2_000_000),
            data_aleatoria(rng, 2025),
        )
        for i in range(10)
    ]
    cab = ["Nº Contrato", "Fornecedor", "CNPJ", "Objeto", "Valor (R$)", "Início da vigência"]

    def escrever(path: Path) -> None:
        aba = Aba("Contratos")
        aba.linha(0, ["Contratos vigentes - Prefeitura Municipal de Santa Luzia do Norte"]).mesclar(
            0, 0, 0, 5
        )
        aba.linha(2, cab)
        for i, (n, f, c, o, v, d) in enumerate(linhas):
            aba.linha(3 + i, [n, f, fmt_cnpj(c), o, v, d])
        escrever_ods(path, [aba])

    saida = tabela(
        ("numero_contrato", TXT, [x[0] for x in linhas]),
        ("fornecedor", TXT, [x[1] for x in linhas]),
        ("cnpj", TXT, [x[2] for x in linhas]),
        ("objeto", TXT, [x[3] for x in linhas]),
        ("valor", NUM, [x[4] for x in linhas]),
        ("inicio_vigencia", DATA, [x[5] for x in linhas]),
    )
    passos = [
        {"op": "select_range", "start": "A3", "end_detect": "auto"},
        {"op": "header_rows", "n": 1},
        {"op": "rename", "mapping": dict(zip(cab, saida.columns, strict=True))},
        {"op": "clean_cnpj", "cols": ["cnpj"]},
        {"op": "cast", "types": {"valor": "float", "inicio_vigencia": "date"}},
    ]
    return Sintetico(
        "transparencia_contratos_ods",
        ("ods",),
        "ods",
        "planilha de contratos publicada em .ods por portal municipal que usa LibreOffice",
        "ODS com título mesclado, CNPJ formatado, valores e datas como tipos nativos do ODS.",
        escrever,
        [Saida(saida, passos, "Contratos")],
    )


@sintetico
def licitacoes_ods() -> Sintetico:
    rng = _rng("licitacoes_ods")
    modalidades = ["Pregão Eletrônico", "Concorrência", "Dispensa", "Tomada de Preços"]
    situacoes = ["Homologada", "Homologada", "Deserta", "Fracassada"]
    linhas = []
    for i in range(10):
        estimado = dinheiro(rng, 10_000, 1_500_000)
        sit = rng.choice(situacoes)
        homologado = round(estimado * rng.uniform(0.7, 1), 2) if sit == "Homologada" else None
        linhas.append(
            (
                f"{1000 + i}/2025",
                rng.choice(modalidades),
                f"Aquisição de material {i + 1}",
                estimado,
                homologado,
                sit,
            )
        )
    cab = ["Processo", "Modalidade", "Objeto", "Valor estimado", "Valor homologado", "Situação"]

    def escrever(path: Path) -> None:
        aba = Aba("Licitações")
        aba.linha(0, ["Município de Santa Luzia do Norte"])
        aba.linha(1, ["Licitações - 2025"])
        aba.linha(3, cab)
        for i, x in enumerate(linhas):
            aba.linha(4 + i, list(x))
        aba.linha(15, ["Fonte: Portal da Transparência Municipal"])
        aba.linha(16, ["Atualizado em 10/01/2026"])
        escrever_ods(path, [aba, Aba("Legenda").linha(0, ["Situação: Homologada, Deserta ou Fracassada"])])

    saida = tabela(
        ("processo", TXT, [x[0] for x in linhas]),
        ("modalidade", TXT, [x[1] for x in linhas]),
        ("objeto", TXT, [x[2] for x in linhas]),
        ("valor_estimado", NUM, [x[3] for x in linhas]),
        ("valor_homologado", NUM, [x[4] for x in linhas]),
        ("situacao", TXT, [x[5] for x in linhas]),
    )
    passos = [
        {"op": "select_range", "start": "A4", "end_detect": "auto"},
        {"op": "header_rows", "n": 1},
        {"op": "rename", "mapping": dict(zip(cab, saida.columns, strict=True))},
        {"op": "cast", "types": {"valor_estimado": "float", "valor_homologado": "float"}},
    ]
    return Sintetico(
        "licitacoes_ods",
        ("ods", "relatorio_rodape"),
        "ods",
        "relatório de licitações em .ods de portal municipal",
        "Duas linhas de título, rodapé com fonte e data, segunda aba de legenda.",
        escrever,
        [Saida(saida, passos, "Licitações")],
    )


# --------------------------------------------------------------------------
# Duas tabelas na mesma aba
# --------------------------------------------------------------------------


@sintetico
def anuario_tabelas_empilhadas() -> Sintetico:
    rng = _rng("anuario_tabelas_empilhadas")
    municipios = [m for _, m in MUNICIPIOS_ES[:8]]
    areas = [round(rng.uniform(80, 3_500), 3) for _ in municipios]
    pops = [rng.randint(10_000, 520_000) for _ in municipios]
    dens = [round(p / a, 2) for p, a in zip(pops, areas, strict=True)]

    def escrever(path: Path) -> None:
        aba = Aba("Território")
        aba.linha(0, ["Tabela 1.1 - Área territorial dos municípios - 2022"], negrito=True)
        aba.linha(1, ["Município", "Área (km²)"], negrito=True)
        for i, (m, a) in enumerate(zip(municipios, areas, strict=True)):
            aba.linha(2 + i, [m, a])
        aba.linha(10, ["Fonte: IBGE - Áreas Territoriais"])
        aba.linha(13, ["Tabela 1.2 - População e densidade demográfica - 2022"], negrito=True)
        aba.linha(14, ["Município", "População", "Densidade (hab/km²)"], negrito=True)
        for i, (m, p, d) in enumerate(zip(municipios, pops, dens, strict=True)):
            aba.linha(15 + i, [m, p, d])
        aba.linha(23, ["Fonte: IBGE - Censo Demográfico 2022"])
        escrever_xlsx(path, [aba])

    t1 = tabela(("municipio", TXT, municipios), ("area_km2", NUM, areas))
    t2 = tabela(("municipio", TXT, municipios), ("populacao", INT, pops), ("densidade", NUM, dens))
    p1 = [
        {"op": "select_range", "start": "A2", "end": "B10"},
        {"op": "header_rows", "n": 1},
        {"op": "rename", "mapping": {"Município": "municipio", "Área (km²)": "area_km2"}},
        {"op": "cast", "types": {"area_km2": "float"}},
    ]
    p2 = [
        {"op": "select_range", "start": "A15", "end": "C23"},
        {"op": "header_rows", "n": 1},
        {
            "op": "rename",
            "mapping": {
                "Município": "municipio",
                "População": "populacao",
                "Densidade (hab/km²)": "densidade",
            },
        },
        {"op": "cast", "types": {"populacao": "int", "densidade": "float"}},
    ]
    return Sintetico(
        "anuario_tabelas_empilhadas",
        ("duas_tabelas",),
        "xlsx",
        "anuário estatístico estadual com várias tabelas por aba",
        "Duas tabelas uma abaixo da outra, cada uma com título e fonte, separadas por linhas em branco.",
        escrever,
        [Saida(t1, p1, "Território"), Saida(t2, p2, "Território")],
    )


@sintetico
def anuario_tabelas_lado_a_lado() -> Sintetico:
    rng = _rng("anuario_tabelas_lado_a_lado")
    municipios = [m for _, m in MUNICIPIOS_ES[4:12]]
    urbana = [rng.randint(1_000, 60_000) for _ in municipios]
    rural = [rng.randint(0, 5_000) for _ in municipios]
    docentes = [rng.randint(50, 4_000) for _ in municipios]
    escolas = [rng.randint(5, 250) for _ in municipios]

    def escrever(path: Path) -> None:
        aba = Aba("Educação")
        aba.linha(0, ["Tabela 2.1 - Matrículas na educação básica - 2024"], negrito=True)
        aba.linha(0, ["Tabela 2.2 - Docentes e escolas - 2024"], c=4, negrito=True)
        aba.linha(1, ["Município", "Urbana", "Rural"], negrito=True)
        aba.linha(1, ["Município", "Docentes", "Escolas"], c=4, negrito=True)
        for i, m in enumerate(municipios):
            aba.linha(2 + i, [m, urbana[i], rural[i]])
            aba.linha(2 + i, [m, docentes[i], escolas[i]], c=4)
        aba.linha(11, ["Fonte: INEP - Censo Escolar da Educação Básica"])
        escrever_xlsx(path, [aba])

    t1 = tabela(("municipio", TXT, municipios), ("urbana", INT, urbana), ("rural", INT, rural))
    t2 = tabela(("municipio", TXT, municipios), ("docentes", INT, docentes), ("escolas", INT, escolas))
    p1 = [
        {"op": "select_range", "start": "A2", "end": "C10"},
        {"op": "header_rows", "n": 1},
        {"op": "rename", "mapping": {"Município": "municipio", "Urbana": "urbana", "Rural": "rural"}},
        {"op": "cast", "types": {"urbana": "int", "rural": "int"}},
    ]
    p2 = [
        {"op": "select_range", "start": "E2", "end": "G10"},
        {"op": "header_rows", "n": 1},
        {"op": "rename", "mapping": {"Município": "municipio", "Docentes": "docentes", "Escolas": "escolas"}},
        {"op": "cast", "types": {"docentes": "int", "escolas": "int"}},
    ]
    return Sintetico(
        "anuario_tabelas_lado_a_lado",
        ("duas_tabelas",),
        "xlsx",
        "anuário estatístico com tabelas lado a lado",
        "Duas tabelas na mesma faixa de linhas, separadas por uma coluna vazia, com fonte comum no fim.",
        escrever,
        [Saida(t1, p1, "Educação"), Saida(t2, p2, "Educação")],
    )


# --------------------------------------------------------------------------
# Mojibake
# --------------------------------------------------------------------------

BAIRROS = ["Jardim Camburi", "Praia do Canto", "Itapoã", "Jardim da Penha", "Centro", "São Pedro", "Maruípe"]
MUNICIPIOS_ACENTO = ["São Mateus", "Conceição da Barra", "Ibiraçu", "João Neiva", "Pedro Canário", "Vitória"]


@sintetico
def cadastro_mojibake() -> Sintetico:
    rng = _rng("cadastro_mojibake")
    linhas = [
        (f"{i + 1:05d}", pessoa(rng), rng.choice(MUNICIPIOS_ACENTO), rng.choice(BAIRROS)) for i in range(12)
    ]
    cab = ["Código", "Nome", "Município", "Bairro"]

    def escrever(path: Path) -> None:
        rows = [[mojibake(c) for c in cab]] + [[mojibake(v) for v in x] for x in linhas]
        escrever_csv(path, rows, encoding="utf-8")

    saida = tabela(
        ("codigo", TXT, [x[0] for x in linhas]),
        ("nome", TXT, [x[1] for x in linhas]),
        ("municipio", TXT, [x[2] for x in linhas]),
        ("bairro", TXT, [x[3] for x in linhas]),
    )
    passos = [
        {"op": "rename", "mapping": {mojibake(c): n for c, n in zip(cab, saida.columns, strict=True)}},
        {"op": "fix_mojibake", "cols": ["nome", "municipio", "bairro"]},
    ]
    return Sintetico(
        "cadastro_mojibake",
        ("mojibake",),
        "csv",
        "cadastro que passou por dois sistemas com encodings diferentes",
        "Arquivo UTF-8 válido, mas com o texto já corrompido ('SÃ£o Mateus'), inclusive no cabeçalho.",
        escrever,
        [Saida(saida, passos)],
    )


@sintetico
def escolas_mojibake() -> Sintetico:
    rng = _rng("escolas_mojibake")
    etapas = ["Educação Infantil", "Ensino Fundamental", "Ensino Médio", "Educação de Jovens e Adultos"]
    linhas = [
        (
            f"EEEF {rng.choice(SOBRENOMES)} {rng.choice(NOMES)}",
            rng.choice(etapas),
            rng.choice(MUNICIPIOS_ACENTO),
        )
        for _ in range(12)
    ]
    cab = ["Escola", "Etapa", "Município"]

    def escrever(path: Path) -> None:
        aba = Aba("Escolas")
        aba.linha(0, cab, negrito=True)
        for i, x in enumerate(linhas):
            # Metade das linhas veio de outro sistema, já corrompida.
            aba.linha(1 + i, [mojibake(v) if i % 2 else v for v in x])
        escrever_xlsx(path, [aba])

    saida = tabela(
        ("escola", TXT, [x[0] for x in linhas]),
        ("etapa", TXT, [x[1] for x in linhas]),
        ("municipio", TXT, [x[2] for x in linhas]),
    )
    passos = [
        {"op": "header_rows", "n": 1},
        {"op": "rename", "mapping": {"Escola": "escola", "Etapa": "etapa", "Município": "municipio"}},
        {"op": "fix_mojibake", "cols": ["escola", "etapa", "municipio"]},
    ]
    return Sintetico(
        "escolas_mojibake",
        ("mojibake",),
        "xlsx",
        "lista de escolas montada a partir de duas fontes",
        "Metade das linhas com acentuação corrompida ('EducaÃ§Ã£o'), metade correta.",
        escrever,
        [Saida(saida, passos, "Escolas")],
    )


# --------------------------------------------------------------------------


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("nomes", nargs="*", help="Arquivos a gerar (padrão: todos).")
    parser.add_argument("--destino", type=Path, default=CORPUS, help="Pasta do corpus.")
    parser.add_argument("--listar", action="store_true", help="Lista os arquivos e as categorias.")
    args = parser.parse_args(argv)

    if args.listar:
        for nome, fn in SINTETICOS.items():
            s = fn()
            print(f"{nome:32} {s.formato:5} {', '.join(s.categorias)}")
        return 0
    desconhecidos = set(args.nomes) - set(SINTETICOS)
    if desconhecidos:
        print(f"Desconhecidos: {', '.join(sorted(desconhecidos))}", file=sys.stderr)
        return 2
    for nome in args.nomes or SINTETICOS:
        pasta = gravar(SINTETICOS[nome](), args.destino)
        print(f"ok  {pasta}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
