"""Exportação da tabela limpa: CSV, Parquet e Excel (F1.I).

- CSV sempre em UTF-8 com BOM (o Excel só reconhece UTF-8 com ele), com
  separador e vírgula decimal opcionais: ``;`` e ``,`` abrem direto no
  Excel em português.
- Excel (``.xlsx``) com os tipos da tabela: data como data, número como
  número. Texto é gravado sempre como texto: um valor ``=SOMA(A1:A9)`` nunca
  vira fórmula no ``.xlsx``.
- Neutralizar fórmulas no CSV (F0.A9): opcional, desligado por padrão.
  Prefixa com ``'`` o texto que começa com ``=``, ``+``, ``-``, ``@``,
  tabulação ou ``\\r`` (recomendação da OWASP). ``celulas_formula`` conta
  quantas células seriam afetadas, para a interface avisar.
- ``gravar_juntos``: build-then-swap. Tudo é gerado numa pasta temporária
  ao lado do destino e só no fim movido para o lugar; uma falha no meio não
  toca em nenhum arquivo existente nem deixa arquivo pela metade (§9).
"""

from __future__ import annotations

import math
import re
import shutil
import tempfile
from collections.abc import Callable, Mapping
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Literal

import polars as pl
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

Separador = Literal[",", ";", "\t"]

# Primeiro caractere que faz o Excel (ou o LibreOffice) tratar a célula de
# um CSV como fórmula.
INICIO_FORMULA = ("=", "+", "-", "@", "\t", "\r")
_RE_FORMULA = r"^[=+\-@\t\r]"

# Limites do formato .xlsx.
MAX_LINHAS_XLSX = 1_048_576  # inclui a linha do cabeçalho
MAX_COLUNAS_XLSX = 16_384
MAX_TEXTO_XLSX = 32_767

_FORMATO_DATA = "DD/MM/YYYY"
_FORMATO_DATA_HORA = "DD/MM/YYYY HH:MM:SS"
_FORMATO_HORA = "HH:MM:SS"


class ExportError(ValueError):
    """A tabela não cabe no formato pedido (linhas demais para o Excel,
    texto longo demais numa célula…). A mensagem é para o usuário."""


@dataclass(frozen=True)
class OpcoesCsv:
    separator: Separador = ","
    decimal_comma: bool = False

    def __post_init__(self) -> None:
        if self.separator not in (",", ";", "\t"):
            raise ValueError(f"separador inválido: {self.separator!r} (use ',', ';' ou tabulação)")


# ---------------------------------------------------------------------------
# Fórmulas (F0.A9)
# ---------------------------------------------------------------------------


def _colunas_texto(df: pl.DataFrame) -> list[str]:
    return [n for n, t in df.schema.items() if t == pl.String]


def celulas_formula(df: pl.DataFrame) -> int:
    """Quantas células de texto começam com um caractere de fórmula."""
    colunas = _colunas_texto(df)
    if not colunas or df.height == 0:
        return 0
    contagem = df.select(pl.col(c).str.contains(_RE_FORMULA).sum() for c in colunas)
    return int(sum(v or 0 for v in contagem.row(0)))


def neutralizar_formulas(df: pl.DataFrame) -> pl.DataFrame:
    """Prefixa com ``'`` o texto que começa com caractere de fórmula. Só
    colunas de texto mudam; número, data e booleano ficam como estão."""
    colunas = _colunas_texto(df)
    if not colunas:
        return df
    return df.with_columns(
        pl.when(pl.col(c).str.contains(_RE_FORMULA)).then("'" + pl.col(c)).otherwise(pl.col(c)).alias(c)
        for c in colunas
    )


# ---------------------------------------------------------------------------
# Escritores (gravam no caminho que recebem)
# ---------------------------------------------------------------------------


def escrever_csv(df: pl.DataFrame, path: Path, opcoes: OpcoesCsv | None = None) -> None:
    """CSV UTF-8 com BOM. Com ``decimal_comma``, número sai como ``1234,5``
    (e entre aspas se o separador também for vírgula). Datas em ISO
    (``2024-01-31``), que o Excel reconhece em qualquer idioma."""
    opcoes = opcoes or OpcoesCsv()
    df.write_csv(
        path,
        include_bom=True,
        separator=opcoes.separator,
        decimal_comma=opcoes.decimal_comma,
    )


def escrever_parquet(df: pl.DataFrame, path: Path) -> None:
    df.write_parquet(path, compression="zstd", statistics=True)


def _nome_aba(nome: str) -> str:
    limpo = re.sub(r"[\[\]:*?/\\]", "_", nome).strip("'") or "dados"
    return limpo[:31]


def _conversor(dtype: pl.DataType) -> tuple[Callable[[Any], Any], str | None]:
    """Função que leva o valor do Polars ao valor da célula, e o formato de
    número do Excel."""
    if dtype == pl.Date:
        return (lambda v: v), _FORMATO_DATA
    if isinstance(dtype, pl.Datetime):
        # O Excel não tem fuso horário: grava a hora local da coluna.
        return (lambda v: v.replace(tzinfo=None) if isinstance(v, datetime) else v), _FORMATO_DATA_HORA
    if dtype == pl.Time:
        return (lambda v: v), _FORMATO_HORA
    if dtype.is_float():
        return (lambda v: None if v is None or math.isnan(v) or math.isinf(v) else v), None
    if dtype.is_numeric() or dtype == pl.Boolean:
        return (lambda v: v), None
    if dtype == pl.String:
        return (lambda v: v), None
    return (lambda v: v if v is None else str(v)), None


def _texto_xlsx(valor: str) -> str:
    return ILLEGAL_CHARACTERS_RE.sub("", valor)  # caracteres de controle que o .xlsx não aceita


def _confere_textos(df: pl.DataFrame) -> None:
    """Texto acima do limite do Excel: erro antes de abrir a planilha (um
    erro no meio da gravação deixaria o arquivo temporário do openpyxl
    aberto)."""
    for nome in _colunas_texto(df):
        tamanhos = df.get_column(nome).str.len_chars()
        maior = tamanhos.max()
        if maior is not None and maior > MAX_TEXTO_XLSX:
            linha = int(tamanhos.arg_max()) + 2  # +1 do cabeçalho, +1 da numeração do Excel
            raise ExportError(
                f"A coluna {nome!r} tem um texto de {maior:,} caracteres na linha {linha}; "
                f"o Excel aceita até {MAX_TEXTO_XLSX:,} por célula. Exporte em CSV ou Parquet."
            )


def escrever_xlsx(df: pl.DataFrame, path: Path, aba: str = "dados") -> None:
    """Excel com os tipos da tabela. Texto é sempre célula de texto (nunca
    fórmula); data, hora e número usam o tipo e o formato do Excel."""
    if df.height + 1 > MAX_LINHAS_XLSX:
        raise ExportError(
            f"A tabela tem {df.height:,} linhas; o Excel aceita até {MAX_LINHAS_XLSX - 1:,} "
            "(fora o cabeçalho). Exporte em CSV ou Parquet."
        )
    if df.width > MAX_COLUNAS_XLSX:
        raise ExportError(f"A tabela tem {df.width:,} colunas; o Excel aceita até {MAX_COLUNAS_XLSX:,}.")
    _confere_textos(df)

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(_nome_aba(aba))
    ws.freeze_panes = "A2"
    for i, nome in enumerate(df.columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(50, max(10, len(nome) + 2))

    negrito = Font(bold=True)
    cabecalho = []
    for nome in df.columns:
        c = WriteOnlyCell(ws, value=_texto_xlsx(nome))
        c.data_type = "s"
        c.font = negrito
        cabecalho.append(c)
    ws.append(cabecalho)

    conversores = [_conversor(t) for t in df.dtypes]
    texto = [t == pl.String for t in df.dtypes]
    for linha in df.iter_rows():
        celulas = []
        for j, valor in enumerate(linha):
            converte, formato = conversores[j]
            v = converte(valor)
            if v is None:
                celulas.append(None)
                continue
            if texto[j]:
                c = WriteOnlyCell(ws, value=_texto_xlsx(v))
                c.data_type = "s"  # nunca fórmula, mesmo começando com "="
            else:
                c = WriteOnlyCell(ws, value=v)
                if formato:
                    c.number_format = formato
            celulas.append(c)
        ws.append(celulas)
    wb.save(path)


# ---------------------------------------------------------------------------
# Build-then-swap (F1.I4)
# ---------------------------------------------------------------------------


def gravar_juntos(escritas: Mapping[Path, Callable[[Path], None]]) -> list[Path]:
    """Grava vários arquivos como uma unidade.

    Cada função recebe um caminho temporário e escreve nele. Só depois que
    todas terminam, os arquivos são movidos para os destinos (renomear na
    mesma pasta é atômico). Se alguma falhar, nenhum destino é tocado e a
    pasta temporária some. Uma função pode gravar uma pasta em vez de um
    arquivo: a pasta antiga do destino é trocada inteira pela nova."""
    por_pasta: dict[Path, list[Path]] = {}
    for destino in escritas:
        por_pasta.setdefault(destino.parent.resolve(), []).append(destino)
    temporarias: dict[Path, Path] = {}
    try:
        for pasta in por_pasta:
            pasta.mkdir(parents=True, exist_ok=True)
            temporarias[pasta] = Path(tempfile.mkdtemp(prefix=".wolfz-export-", dir=pasta))
        prontos: list[tuple[Path, Path]] = []
        for destino, escrever in escritas.items():
            tmp = temporarias[destino.parent.resolve()] / destino.name
            escrever(tmp)
            prontos.append((tmp, destino))
        for tmp, destino in prontos:
            if tmp.is_dir() and destino.is_dir():
                # Pasta (o PBIP): a antiga vai para a temporária, que some no fim.
                destino.replace(tmp.with_name(tmp.name + ".antiga"))
            tmp.replace(destino)
    finally:
        for tmp in temporarias.values():
            shutil.rmtree(tmp, ignore_errors=True)
    return list(escritas)


EXTENSOES = {"csv": ".csv", "parquet": ".parquet", "xlsx": ".xlsx"}


def escritor(
    formato: str,
    df: pl.DataFrame,
    csv: OpcoesCsv | None = None,
    neutralizar: bool = False,
    aba: str = "dados",
) -> Callable[[Path], None]:
    """Função que grava ``df`` no formato pedido (para ``gravar_juntos``).

    ``neutralizar`` só muda o CSV: no ``.xlsx`` o texto já é sempre célula de
    texto (o apóstrofo apareceria no valor), e o Parquet não é aberto por
    planilha."""
    if neutralizar and formato == "csv":
        df = neutralizar_formulas(df)
    if formato == "csv":
        return lambda p: escrever_csv(df, p, csv)
    if formato == "parquet":
        return lambda p: escrever_parquet(df, p)
    if formato == "xlsx":
        return lambda p: escrever_xlsx(df, p, aba)
    raise ValueError(f"formato de dados desconhecido: {formato!r}")


__all__ = [
    "EXTENSOES",
    "INICIO_FORMULA",
    "MAX_LINHAS_XLSX",
    "ExportError",
    "OpcoesCsv",
    "celulas_formula",
    "escrever_csv",
    "escrever_parquet",
    "escrever_xlsx",
    "escritor",
    "gravar_juntos",
    "neutralizar_formulas",
]
