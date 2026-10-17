"""Leitura de planilhas: abas e valores como texto, na posição original.

- ``.xlsx``: valores com ``fastexcel`` (Rust, direto para Arrow).
- ``.xls`` e ``.ods``: valores com ``python-calamine``.
- Estrutura (mesclagens, negrito, fórmulas) em ``structure.py``.

A tabela devolvida é a grade da aba, sem cabeçalho: a linha ``i`` do
DataFrame é a linha ``i + 1`` da planilha e cada coluna tem o nome da letra
da coluna do Excel (``A``, ``B``, …, ``AA``). Colunas vazias à esquerda da
primeira célula usada não aparecem; o nome da primeira coluna diz onde a
grade começa. Decidir o que é cabeçalho, dado e total é da detecção.
"""

from __future__ import annotations

import zipfile
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import asdict, dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Literal

import fastexcel
import openpyxl
import polars as pl
import python_calamine
from python_calamine import CalamineWorkbook, SheetTypeEnum, SheetVisibleEnum

from wolfz.core.ingest import errors
from wolfz.core.ingest.cells import column_letter
from wolfz.core.ingest.errors import IngestError
from wolfz.core.ingest.formats import Format, detect_format
from wolfz.core.ingest.limits import DEFAULT_LIMITS, Limits, check_sheet_count
from wolfz.core.ingest.structure import (
    HEAD_ROWS,
    TAIL_ROWS,
    SheetStructure,
    calamine_structure,
    read_xlsx_structure,
    sample_windows,
)

Visibility = Literal["visible", "hidden", "very_hidden"]


def _visibility(value: Any) -> Visibility:
    # Os enums do python-calamine não são hasheáveis; comparação direta.
    if value == SheetVisibleEnum.Hidden:
        return "hidden"
    if value == SheetVisibleEnum.VeryHidden:
        return "very_hidden"
    return "visible"


@dataclass(frozen=True)
class SheetInfo:
    """Aba de uma planilha. ``rows``/``columns`` são estimativas (no
    ``.xlsx``, a dimensão que o próprio arquivo declara) e podem ser
    ``None`` quando o arquivo não informa."""

    index: int
    name: str
    visibility: Visibility
    rows: int | None
    columns: int | None

    @property
    def hidden(self) -> bool:
        return self.visibility != "visible"

    def to_dict(self) -> dict[str, Any]:
        return asdict(self) | {"hidden": self.hidden}


@dataclass
class SheetData:
    sheet: str
    frame: pl.DataFrame
    structure: SheetStructure


@contextmanager
def reading(path: Path) -> Iterator[None]:
    """Traduz as exceções das bibliotecas de leitura em ``IngestError``."""
    try:
        yield
    except IngestError:
        raise
    except PermissionError as exc:
        raise errors.locked(path.name) from exc
    except python_calamine.PasswordError as exc:
        raise errors.password_protected(path.name) from exc
    except (zipfile.BadZipFile, python_calamine.CalamineError, fastexcel.FastExcelError) as exc:
        raise errors.corrupted(path.name, _first_line(exc)) from exc


def _first_line(exc: Exception) -> str:
    text = str(exc).strip()
    return text.splitlines()[0] if text else type(exc).__name__


def _worksheets(wb: CalamineWorkbook) -> list[Any]:
    return [m for m in wb.sheets_metadata if m.typ == SheetTypeEnum.WorkSheet]


def _xlsx_dimensions(path: Path) -> dict[str, tuple[int | None, int | None]]:
    """Dimensão declarada por aba (tag ``<dimension>``), sem ler as células.
    Arquivo que o ``openpyxl`` não entende fica sem dimensão (§10.5)."""
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
    except Exception:
        return {}
    try:
        return {ws.title: (ws.max_row, ws.max_column) for ws in wb.worksheets}
    except Exception:
        return {}
    finally:
        wb.close()


def list_sheets(
    path: str | Path, fmt: Format | None = None, limits: Limits = DEFAULT_LIMITS, dimensions: bool = True
) -> list[SheetInfo]:
    """Abas de planilha (gráficos ficam de fora), com visibilidade e
    dimensão estimada. ``dimensions=False`` pula a dimensão, que no
    ``.xls``/``.ods`` exige ler cada aba."""
    path = Path(path)
    fmt = fmt or detect_format(path)
    with reading(path), CalamineWorkbook.from_path(str(path)) as wb:
        metas = _worksheets(wb)
        check_sheet_count(path.name, len(metas), limits)
        dims: dict[str, tuple[int | None, int | None]] = {}
        if dimensions and fmt == "xlsx":
            dims = _xlsx_dimensions(path)
        elif dimensions:
            for m in metas:
                end = wb.get_sheet_by_name(m.name).end
                dims[m.name] = (end[0] + 1, end[1] + 1) if end else (0, 0)
    return [
        SheetInfo(i, m.name, _visibility(m.visible), *dims.get(m.name, (None, None)))
        for i, m in enumerate(metas)
    ]


def resolve_sheet(path: Path, sheets: list[SheetInfo], sheet: str | int | None) -> str:
    """Nome da aba pedida. Sem aba, a primeira visível."""
    names = [s.name for s in sheets]
    if sheet is None:
        visible = [s.name for s in sheets if not s.hidden]
        if not visible and not names:
            raise IngestError("no_sheets", "Planilha sem abas", f"{path.name} não tem nenhuma aba de dados.")
        return (visible or names)[0]
    if isinstance(sheet, int):
        if 0 <= sheet < len(names):
            return names[sheet]
    elif sheet in names:
        return sheet
    raise errors.sheet_not_found(path.name, sheet, names)


def _text(value: Any) -> str | None:
    """Valor do ``python-calamine`` como texto, no mesmo formato que o
    ``fastexcel`` produz para o ``.xlsx`` (receita feita num formato
    funciona no outro)."""
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() and abs(value) < 1e21 else repr(value)
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d 00:00:00")
    if isinstance(value, time | timedelta):
        return str(value)
    return str(value)


def _read_xlsx(path: Path, sheet: str) -> pl.DataFrame:
    reader = fastexcel.read_excel(str(path))
    # skip_rows=0 com header_row=None mantém as linhas vazias do topo:
    # a linha i do DataFrame é a linha i da planilha.
    loaded = reader.load_sheet(sheet, header_row=None, dtypes="string", skip_rows=0)
    columns = loaded.available_columns()
    df = loaded.to_polars()
    if not columns or df.width == 0:
        return pl.DataFrame()
    return df.rename({c.name: column_letter(c.absolute_index) for c in columns})


def _read_calamine(wb: CalamineWorkbook, sheet: str) -> tuple[pl.DataFrame, list[list[Any]], Any]:
    ws = wb.get_sheet_by_name(sheet)
    grid = ws.to_python(skip_empty_area=False)
    try:
        merged = ws.merged_cell_ranges
    except python_calamine.CalamineError:
        merged = None
    if not ws.start:
        return pl.DataFrame(), grid, merged
    first_col = ws.start[1]
    width = max((len(r) for r in grid), default=0)
    data = {
        column_letter(c): [_text(r[c]) if c < len(r) else None for r in grid] for c in range(first_col, width)
    }
    return pl.DataFrame(data, schema=dict.fromkeys(data, pl.String)), grid, merged


def read_sheet(
    path: str | Path,
    sheet: str | int | None = None,
    fmt: Format | None = None,
    head_rows: int = HEAD_ROWS,
    tail_rows: int = TAIL_ROWS,
    sheets: list[SheetInfo] | None = None,
) -> SheetData:
    """Lê uma aba: valores (todos como texto) e estrutura das primeiras e
    últimas linhas."""
    path = Path(path)
    fmt = fmt or detect_format(path)
    if fmt == "csv":
        raise ValueError("read_sheet lê planilhas; para CSV use wolfz.core.ingest.csv")
    sheets = sheets if sheets is not None else list_sheets(path, fmt, dimensions=False)
    name = resolve_sheet(path, sheets, sheet)
    if fmt == "xlsx":
        with reading(path):
            df = _read_xlsx(path, name)
        try:
            structure = read_xlsx_structure(path, name, df.height, head_rows, tail_rows)
        except Exception:
            # §10.5: sem formatação, a detecção segue só com os valores.
            structure = SheetStructure(
                name,
                sampled_rows=sample_windows(df.height, head_rows, tail_rows),
                warnings=["Não foi possível ler a formatação deste arquivo; a detecção usa só os valores."],
            )
        return SheetData(name, df, structure)
    with reading(path), CalamineWorkbook.from_path(str(path)) as wb:
        df, grid, merged = _read_calamine(wb, name)
    return SheetData(name, df, calamine_structure(name, grid, merged, fmt, head_rows, tail_rows))


def read_sheet_raw(path: str | Path, sheet: str | int | None = None) -> pl.DataFrame:
    """Só os valores de uma aba, como texto, na posição original (sem ler a
    estrutura)."""
    path = Path(path)
    fmt = detect_format(path)
    if fmt == "csv":
        raise ValueError("read_sheet_raw lê planilhas; para CSV use wolfz.core.ingest.csv")
    name = resolve_sheet(path, list_sheets(path, fmt, dimensions=False), sheet)
    with reading(path):
        if fmt == "xlsx":
            return _read_xlsx(path, name)
        with CalamineWorkbook.from_path(str(path)) as wb:
            return _read_calamine(wb, name)[0]
